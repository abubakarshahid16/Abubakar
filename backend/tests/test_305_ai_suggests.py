"""Order 305.4 part B: the CRS column "AI suggests (internal)": Pass / Fail / Query with the value, the rule and the
clause, beside the row and never in the verdict. Internal copy only. Invented findings only."""

from __future__ import annotations

import io

import openpyxl

from app import crs_export, crs_mapping, crs_suggestion

HELD = {
    "id": "f1",
    "origin": "datasheet_check",
    "compliance_status": "NEEDS_ENGINEER_REVIEW",
    "rule_said": "NON_COMPLIANT",
    "trust_label": "UNPROVEN",
    "contractor_section": "Set pressure",
    "contractor_evidence_text": "12 barg",
    "contractor_page": 2,
    "requirement_source_text": "Datasheet check PSV-24: The set pressure is at most the MAWP.",
    "ai_rationale": "Datasheet check PSV-24: 12 barg is above the MAWP of 10 barg (the rule is not yet proven)",
    "finding": "Datasheet check: held",
    "required_action": "Engineer to check.",
}


def test_a_held_rule_that_said_does_not_meet_is_a_fail_suggestion():
    text = crs_suggestion.words(HELD, "API 520-I 5.4.2.1")
    assert text.startswith("Fail (sheet: Set pressure = 12 barg; rule: ")
    assert "clause: API 520-I 5.4.2.1)" in text and crs_suggestion.NOTE in text


def test_a_held_rule_that_said_meets_is_a_pass_suggestion():
    assert crs_suggestion.words({**HELD, "rule_said": "COMPLIANT"}, "").startswith(
        "Pass ("
    )


def test_a_question_for_the_engineer_with_no_suggestion_is_a_query():
    query = {**HELD, "rule_said": None}
    assert crs_suggestion.words(query, "").startswith("Query (")


def test_a_verdict_code_gave_a_missing_value_and_a_bare_finding_suggest_nothing():
    for status in (
        "COMPLIANT",
        "NON_COMPLIANT",
        "MISSING_INFORMATION",
        "NOT_APPLICABLE",
        None,
    ):
        assert (
            crs_suggestion.words(
                {**HELD, "compliance_status": status, "rule_said": None}, "x"
            )
            == ""
        )
    assert crs_suggestion.words({"id": "x"}, "") == ""


def test_a_suggestion_with_nothing_to_show_is_just_the_word_and_the_note():
    bare = {"compliance_status": "NEEDS_ENGINEER_REVIEW"}
    assert crs_suggestion.words(bare, "") == f"Query - {crs_suggestion.NOTE}"


def test_a_group_prints_each_distinct_suggestion_once():
    twin = {**HELD, "id": "f2"}
    other = {**HELD, "id": "f3", "contractor_evidence_text": "13 barg"}
    out = crs_suggestion.for_group([HELD, twin, other], lambda f: "")
    assert len(out.splitlines()) == 2


def _rows():
    return crs_mapping.build_crs_rows([HELD], [], "sheet.pdf")


def test_the_row_carries_the_suggestion_and_the_verdict_cells_do_not_change():
    row = next(r for r in _rows() if r.get("ai_suggests"))
    assert row["ai_suggests"].startswith("Fail (")
    other = next(
        r
        for r in crs_mapping.build_crs_rows(
            [{**HELD, "rule_said": "COMPLIANT"}], [], "sheet.pdf"
        )
        if r.get("ai_suggests")
    )
    assert other["ai_suggests"].startswith("Pass (") and row["ai_suggests"].startswith(
        "Fail ("
    )
    for key in (
        "comment",
        "priority",
        "severity",
        "row_kind",
        "standard_reference",
        "page_section",
    ):
        assert row[key] == other[key], key  # the suggestion is its own cell


def _book(copy):
    rows = _rows()
    for r in rows:
        r["engineer_confirmed"] = True
    meta = {"review_run_id": "run-1", "copy": copy}
    view = crs_export.build_crs_view(rows, meta)
    ws = openpyxl.load_workbook(io.BytesIO(crs_export.build_crs(rows, meta))).active
    return view, ws


def test_the_internal_copy_has_the_column_last_and_prints_it():
    view, ws = _book("internal")
    assert view["columns"][-1] == crs_suggestion.COLUMN
    assert (
        ws.cell(row=crs_export.COLUMN_HEADER_ROW, column=13).value
        == crs_suggestion.COLUMN
    )
    assert str(
        ws.cell(row=crs_export.COLUMN_HEADER_ROW + 1, column=13).value
    ).startswith("Fail (")


def test_the_contractor_copy_never_carries_the_column_or_the_text():
    view, ws = _book("issue")
    assert crs_suggestion.COLUMN not in view["columns"]
    assert all(r["ai_suggests"] == "" for r in view["rows"])
    cells = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert not any("suggestion only" in str(v) for v in cells)
