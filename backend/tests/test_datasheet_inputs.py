"""Datasheet inputs other than a PDF text layer: Excel, Word, scanned pages.

Every file here is MADE UP and generated in the test (tmp_path); nothing is
committed as a binary. Values use the fixture conventions (DS-/EX-/SYN-,
0000/1234/9999).

Behind DATASHEET_OFFICE_INPUT, off by default: the last section proves that
with the flag off nothing changes.

Mutations: M1700-M1719 (`scripts/mutations/datasheet_inputs.py`).
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from app import datasheet_inputs as di
from app import datasheets, db, states, submittal_review, upload
from app.config import settings

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


@pytest.fixture(autouse=True)
def temp_storage(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "db_path", tmp_path / "dsin.sqlite")
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    (tmp_path / "uploads").mkdir(parents=True, exist_ok=True)
    db.reset_connection(); db.init_db(); submittal_review.ensure_schema()
    submittal_review.migrate_facts_to_per_document()
    di._cached_office.cache_clear(); di._kind_cached.cache_clear()
    yield
    db.reset_connection()


@pytest.fixture
def flag_on(monkeypatch):
    monkeypatch.setattr(settings, "datasheet_office_input", True)


# ------------------------------------------------------------------ builders


def _xlsx(path: Path, build) -> str:
    import openpyxl
    wb = openpyxl.Workbook()
    build(wb)
    wb.save(str(path))
    return str(path)


def _datasheet_book(wb):
    ws = wb.active
    ws.title = "DS-0001"
    ws.append(["Item", "Value", "Unit"])
    ws.append([1, "Design pressure", 15, "barg"])
    ws.append(["Design temperature", 120.5, "degC"])
    ws.append(["Tag number", "EX-1234"])


def _docx_bytes(body: str, *, prolog: str = "", extra: dict[str, str] | None = None) -> bytes:
    xml = (f'<?xml version="1.0" encoding="UTF-8"?>{prolog}'
           f'<w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", xml)
        for name, text in (extra or {}).items():
            z.writestr(name, text)
    return buf.getvalue()


def _p(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def _tc(text: str, props: str = "") -> str:
    return f"<w:tc>{f'<w:tcPr>{props}</w:tcPr>' if props else ''}{_p(text)}</w:tc>"


def _tbl(*rows: str) -> str:
    return "<w:tbl>" + "".join(f"<w:tr>{r}</w:tr>" for r in rows) + "</w:tbl>"


def _datasheet_docx(path: Path) -> str:
    body = (_p("SYN-0000 valve datasheet")
            + _tbl(_tc("Design pressure") + _tc("15") + _tc("barg"),
                   _tc("Design temperature") + _tc("120") + _tc("degC"))
            + _p("Body material: A105"))
    path.write_bytes(_docx_bytes(body))
    return str(path)


def _ingest_office(path: str, doc_id: str) -> str:
    """Document + one chunk per page, the chunk text being the page text -
    what the extract stage stores and the chunker cites."""
    pages = di.page_texts(path)
    with db.connect() as conn:
        conn.execute("""INSERT INTO documents
            (id,filename,sha256,size_bytes,stored_path,status,page_count,uploaded_at)
            VALUES (?,?,?,1,?,'ready',?,?)""",
            (doc_id, f"{doc_id}.bin", f"sha-{doc_id}", path, len(pages),
             "2026-09-30T00:00:00Z"))
        for page in pages:
            conn.execute("""INSERT INTO chunks
                (id,document_id,filename,ordinal,page_start,page_end,section,kind,
                 text,token_count,content_hash,retrievable)
                VALUES (?,?,?,?,?,?,NULL,'prose',?,1,?,1)""",
                (f"{doc_id}-c{page.page_no}", doc_id, f"{doc_id}.bin", page.page_no,
                 page.page_no, page.page_no, page.text, f"h{doc_id}{page.page_no}"))
    return doc_id


def _facts(doc_id: str) -> list[dict]:
    return datasheets.list_facts(doc_id, allowed_document_ids=frozenset({doc_id}))


def _extract(doc_id: str) -> dict:
    return datasheets.extract_facts(doc_id, allowed_document_ids=frozenset({doc_id}))


# ------------------------------------------------------------------ xlsx


def test_xlsx_rows_become_facts_provable_against_the_page_text(tmp_path, flag_on):
    path = _xlsx(tmp_path / "ds.xlsx", _datasheet_book)
    doc = _ingest_office(path, "doc_xlsx")
    out = _extract(doc)
    assert out["facts"] > 0, out

    by = {f["field_name"]: f for f in _facts(doc)}
    pressure = by["design pressure"]
    assert pressure["raw_value"] == "15" and pressure["raw_unit"] == "barg"
    assert pressure["extraction_method"] == "xlsx"
    assert by["design temperature"]["raw_value"] == "120.5"
    # PROVABLE: the cited chunk holds the row, label and value cells verbatim.
    chunk = db.connect().execute("SELECT text FROM chunks WHERE id = ?",
                                 (pressure["chunk_id"],)).fetchone()["text"]
    assert "Design pressure | 15 | barg" in chunk
    # the header row "Item | Value | Unit" is not a fact
    assert "item" not in by


def test_every_xlsx_pair_is_found_on_its_page_line(tmp_path):
    path = _xlsx(tmp_path / "ds.xlsx", _datasheet_book)
    page = di.page_texts(path)[0]
    assert page.source == "xlsx"
    lines = page.text.splitlines()
    for row in page.rows:
        assert di.CELL_SEPARATOR.join(row) in lines
    pairs = di.pairs_from_rows(page.rows)
    assert ("Design pressure", "15 barg") in pairs
    for label, value in pairs:
        line = next(line for line in lines if label in line.split(di.CELL_SEPARATOR))
        assert all(cell in line for cell in value.split())


def test_merged_cells_label_every_row_they_span(tmp_path):
    def build(wb):
        ws = wb.active
        ws["A1"] = "Body material"
        ws.merge_cells("A1:A2")
        ws["B1"] = "A105"
        ws["B2"] = "A350 LF2"
        ws["A4"] = "Design code"
        ws.merge_cells("A4:B4")
        ws["C4"] = "EX-9999"
    rows = di.page_texts(_xlsx(tmp_path / "m.xlsx", build))[0].rows
    assert ("Body material", "A105") in rows
    assert ("Body material", "A350 LF2") in rows
    # a horizontal merge is written once
    assert ("Design code", "EX-9999") in rows


def test_formula_cells_are_never_evaluated_or_injected(tmp_path):
    def build(wb):
        ws = wb.active
        ws["A1"] = "Design pressure"
        ws["B1"] = "=1234+1"
        ws["C1"] = "barg"
    page = di.page_texts(_xlsx(tmp_path / "f.xlsx", build))[0]
    assert "=" not in page.text and "1234" not in page.text and "1235" not in page.text
    assert ("Design pressure", "barg") in page.rows


def test_hidden_sheet_is_skipped_and_reported(tmp_path):
    def build(wb):
        wb.active.append(["Design pressure", 15, "barg"])
        hidden = wb.create_sheet("SYN-hidden")
        hidden.append(["Poison field", 9999, "barg"])
        hidden.sheet_state = "hidden"
    reading = di.read(_xlsx(tmp_path / "h.xlsx", build))
    assert len(reading.pages) == 1
    assert "Poison field" not in reading.pages[0].text
    assert "hidden sheet 2 skipped" in reading.notes


def test_a_sheet_past_the_row_limit_is_cut_and_says_so(tmp_path, monkeypatch):
    monkeypatch.setattr(di, "MAX_SHEET_ROWS", 3)

    def build(wb):
        for i in range(10):
            wb.active.append([f"Field {i}", i])
    reading = di.read(_xlsx(tmp_path / "big.xlsx", build))
    assert len(reading.pages[0].rows) == 3
    assert "sheet 1 cut at 3 rows" in reading.notes


# ------------------------------------------------------------------ docx


def test_docx_table_becomes_facts(tmp_path, flag_on):
    path = _datasheet_docx(tmp_path / "ds.docx")
    doc = _ingest_office(path, "doc_docx")
    assert _extract(doc)["facts"] > 0
    by = {f["field_name"]: f for f in _facts(doc)}
    assert by["design pressure"]["raw_value"] == "15"
    assert by["design pressure"]["raw_unit"] == "barg"
    assert by["design pressure"]["extraction_method"] == "docx"
    assert by["body material"]["field_value"] == "A105"


def test_docx_keeps_document_order_and_vertical_merges(tmp_path):
    body = (_p("First paragraph")
            + _tbl(_tc("Body material", '<w:vMerge w:val="restart"/>') + _tc("A105"),
                   _tc("", "<w:vMerge/>") + _tc("A350 LF2"))
            + _p("Last paragraph"))
    (tmp_path / "v.docx").write_bytes(_docx_bytes(body))
    page = di.page_texts(tmp_path / "v.docx")[0]
    assert page.text.splitlines() == [
        "First paragraph", "Body material | A105", "Body material | A350 LF2",
        "Last paragraph"]


def test_docx_with_a_doctype_is_refused(tmp_path, flag_on):
    prolog = '<!DOCTYPE d [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
    (tmp_path / "x.docx").write_bytes(_docx_bytes(_p("Design pressure: &x;"), prolog=prolog))
    with pytest.raises(di.InputError) as err:
        di.read(tmp_path / "x.docx")
    assert err.value.code == "doctype"
    # and the datasheet reader names the FILE's condition, reading nothing
    doc = _ingest_office_unreadable(str(tmp_path / "x.docx"), "doc_xxe")
    out = _extract(doc)
    assert out["facts"] == 0 and out["pages_unreadable"] == 1
    assert out["unreadable"][0]["rule"] == "office_doctype"


def _ingest_office_unreadable(path: str, doc_id: str) -> str:
    with db.connect() as conn:
        conn.execute("""INSERT INTO documents
            (id,filename,sha256,size_bytes,stored_path,status,page_count,uploaded_at)
            VALUES (?,?,?,1,?,'ready',1,'2026-09-30T00:00:00Z')""",
            (doc_id, f"{doc_id}.bin", f"sha-{doc_id}", path))
        conn.execute("""INSERT INTO chunks
            (id,document_id,filename,ordinal,page_start,page_end,section,kind,
             text,token_count,content_hash,retrievable)
            VALUES (?,?,?,1,1,1,NULL,'prose','x',1,?,1)""",
            (f"{doc_id}-c1", doc_id, f"{doc_id}.bin", f"h{doc_id}"))
    return doc_id


def test_a_doctype_in_any_xml_part_is_refused(tmp_path):
    """Not only word/document.xml: openpyxl and any later reader parse the
    other parts, so a DOCTYPE anywhere refuses the file."""
    (tmp_path / "s.docx").write_bytes(_docx_bytes(
        _p("Design pressure: 15 barg"),
        extra={"word/styles.xml": '<!DOCTYPE s [<!ENTITY a "aaaa">]><styles/>'}))
    with pytest.raises(di.InputError) as err:
        di.read(tmp_path / "s.docx")
    assert err.value.code == "doctype"


def test_the_document_part_parser_refuses_a_doctype_on_its_own():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", '<!DOCTYPE d [<!ENTITY a "b">]><d>&a;</d>')
    with zipfile.ZipFile(buf) as z, pytest.raises(di.InputError) as err:
        di._parse_bounded(z, "word/document.xml")
    assert err.value.code == "doctype"


def test_an_oversized_zip_is_refused_before_it_is_expanded(tmp_path, monkeypatch):
    monkeypatch.setattr(di, "MAX_UNCOMPRESSED_BYTES", 50_000)
    body = _p("Design pressure: 15 barg") + _p("0" * 100_000)
    (tmp_path / "bomb.docx").write_bytes(_docx_bytes(body))
    assert (tmp_path / "bomb.docx").stat().st_size < 50_000  # small on disk
    with pytest.raises(di.InputError) as err:
        di.read(tmp_path / "bomb.docx")
    assert err.value.code == "too_large_unpacked"


# ------------------------------------------------------------------ OCR


def _scanned(tmp_path, doc_id: str, ocr_text: str) -> str:
    import pymupdf
    doc = pymupdf.open()
    doc.new_page(width=600, height=500)
    path = tmp_path / f"{doc_id}.pdf"
    doc.save(str(path)); doc.close()
    with db.connect() as conn:
        conn.execute("""INSERT INTO documents
            (id,filename,sha256,size_bytes,stored_path,status,page_count,uploaded_at)
            VALUES (?,?,?,1,?,'ready',1,'2026-09-30T00:00:00Z')""",
            (doc_id, "scan.pdf", f"sha-{doc_id}", str(path)))
        conn.execute("""INSERT INTO chunks
            (id,document_id,filename,ordinal,page_start,page_end,section,kind,
             text,token_count,content_hash,retrievable)
            VALUES (?,?,'scan.pdf',1,1,1,NULL,'prose',' ',1,?,1)""",
            (f"{doc_id}-c1", doc_id, f"h{doc_id}"))
        conn.execute(
            """INSERT INTO page_ocr
               (document_id, page_no, text, char_count, engine, model, dpi,
                mean_conf, box_count, seconds, recognised_at, batch_no)
               VALUES (?,1,?,?, 'rapidocr-3.9.2', 'test-model', 200, 0.81, 1,
                       0.1, '2026-09-30T00:00:00Z', 0)""",
            (doc_id, ocr_text, len(ocr_text)))
    return str(path)


OCR_TABLE = "SYN-0000 DATASHEET\nDesign pressure    15    barg\nBody material | A105"


def test_a_scanned_table_line_becomes_a_low_confidence_flagged_fact(tmp_path, flag_on):
    _scanned(tmp_path, "doc_scan", OCR_TABLE)
    assert _extract("doc_scan")["facts"] >= 2
    by = {f["field_name"]: f for f in _facts("doc_scan")}
    pressure = by["design pressure"]
    assert pressure["raw_value"] == "15" and pressure["raw_unit"] == "barg"
    assert pressure["extraction_method"] == "ocr_fallback"
    assert pressure["confidence"] < datasheets.LOW_CONFIDENCE_THRESHOLD
    assert pressure["validation_state"] == datasheets.NEEDS_ENGINEER_REVIEW
    assert by["body material"]["field_value"] == "A105"


def test_a_scanned_pdf_page_reads_as_ocr_text_with_its_confidence(tmp_path):
    path = _scanned(tmp_path, "doc_scan2", OCR_TABLE)
    page = di.page_texts(path, document_id="doc_scan2")[0]
    assert page.source == "ocr" and page.confidence == pytest.approx(0.81)
    assert "Design pressure" in page.text
    assert ("Design pressure", "15", "barg") in page.rows
    # without the document's stored OCR (and no recognition asked for), the
    # page is honestly empty - nothing invented
    empty = di.page_texts(path)[0]
    assert empty.source == "pdf_text" and not empty.text.strip()


# ------------------------------------------------------------------ wiring


def test_docx_upload_is_accepted_and_indexed_with_the_flag_on(flag_on):
    row, job_id, _dup = upload.ingest(
        io.BytesIO(_docx_bytes(_p("Design pressure: 15 barg"))), "DS-0001.docx")
    assert row["stored_path"].endswith(".docx")
    assert row["filename"] == "DS-0001.docx"
    assert row["status"] == states.QUEUED and job_id


def test_docx_upload_bomb_is_refused(flag_on, monkeypatch):
    monkeypatch.setattr(upload, "MAX_XLSX_UNCOMPRESSED_BYTES", 50_000)
    data = _docx_bytes(_p("0" * 100_000))
    with pytest.raises(upload.UploadError) as err:
        upload.ingest(io.BytesIO(data), "DS-0002.docx")
    assert err.value.code == "not_docx"
    assert not list(Path(settings.upload_dir).glob("*.docx"))


def test_the_extract_stage_writes_office_pages(tmp_path, flag_on):
    from app import extract
    path = _xlsx(tmp_path / "e.xlsx", _datasheet_book)
    with db.connect() as conn:
        conn.execute("""INSERT INTO documents
            (id,filename,sha256,size_bytes,stored_path,status,uploaded_at)
            VALUES ('doc_e','e.xlsx','sha-e',1,?,'extracting','2026-09-30T00:00:00Z')""",
            (path,))
        conn.execute("""INSERT INTO jobs (id, document_id, stage, state, started_at, updated_at)
            VALUES ('job_e','doc_e','extract','running','2026-09-30T00:00:00Z',
                    '2026-09-30T00:00:00Z')""")
    out = extract.extract_document("doc_e")
    assert out["pages_extracted"] == 1
    text = db.connect().execute(
        "SELECT text FROM pages WHERE document_id='doc_e' AND page_no=1").fetchone()["text"]
    assert "Design pressure | 15 | barg" in text


# ------------------------------------------------------------------ flag off


def test_flag_off_docx_upload_is_refused_as_before():
    assert settings.datasheet_office_input is False
    with pytest.raises(upload.UploadError) as err:
        upload.ingest(io.BytesIO(_docx_bytes(_p("x"))), "DS-0003.docx")
    assert err.value.code == "not_xlsx"


def test_flag_off_xlsx_is_stored_not_indexed_and_reads_no_facts(tmp_path):
    buf = io.BytesIO()
    import openpyxl
    wb = openpyxl.Workbook(); _datasheet_book(wb); wb.save(buf)
    row, job_id, _dup = upload.ingest(io.BytesIO(buf.getvalue()), "DS-0004.xlsx")
    assert row["status"] == states.STORED_NOT_INDEXED and job_id is None
    # even a workbook that somehow has chunks is not read by the office path
    path = _xlsx(tmp_path / "off.xlsx", _datasheet_book)
    doc = _ingest_office(path, "doc_off")
    _extract(doc)
    assert not [f for f in _facts(doc) if f["extraction_method"] in ("xlsx", "docx")]
    assert "design pressure" not in {f["field_name"] for f in _facts(doc)}


def test_flag_off_ocr_table_lines_are_not_read(tmp_path):
    _scanned(tmp_path, "doc_scan_off", OCR_TABLE)
    _extract("doc_scan_off")
    assert "design pressure" not in {f["field_name"] for f in _facts("doc_scan_off")}


def test_an_uploaded_workbook_goes_through_the_whole_pipeline_into_facts(flag_on):
    """Upload -> extract stage -> chunker -> keyword index -> datasheet reader.

    The chunker's quality gate drops bare label/value lines as prose without
    a clause, so this only works because the sheet's rows are stored as the
    page's table (`datasheet_inputs.tables_json`) and kept whole."""
    import openpyxl

    from app import ingest

    wb = openpyxl.Workbook()
    _datasheet_book(wb)
    for label, value, unit in [("Design flow", 120, "m3/h"), ("Operating pressure", 9, "barg"),
                               ("Operating temperature", 60, "degC"),
                               ("Corrosion allowance", 3, "mm"), ("Test pressure", 23, "barg"),
                               ("Density", 850, "kg/m3"), ("Viscosity", 2, "cP")]:
        wb.active.append([label, value, unit])
    buf = io.BytesIO(); wb.save(buf)
    row, job_id, _dup = upload.ingest(io.BytesIO(buf.getvalue()), "DS-0005.xlsx")
    assert job_id, "the workbook was not queued for indexing"
    ingest.IngestionWorker().process(row["id"])
    chunk = db.connect().execute(
        "SELECT kind, text FROM chunks WHERE document_id = ? AND retrievable = 1",
        (row["id"],)).fetchone()
    assert chunk is not None and chunk["kind"] == "table"

    assert _extract(row["id"])["facts"] >= 8
    by = {f["field_name"]: f for f in _facts(row["id"])}
    assert by["design pressure"]["raw_value"] == "15"
    assert by["density"]["raw_unit"] == "kg/m3"
    assert {f["extraction_method"] for f in by.values()} == {"xlsx"}
