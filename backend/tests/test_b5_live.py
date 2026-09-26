"""Master order B5 on the LIVE review path (2026-09-25).

What changed, and what each test holds:

  1. SELECTION ON EVIDENCE. A standard is applied when the submittal cites it,
     when it is classified for the submittal's equipment, service or project,
     or when its own scope clause names the equipment. A shared discipline or
     similar wording alone is recorded as CONSIDERED, NOT APPLIED, with the
     reason - it used to be applied, so every mechanical standard was compared
     against every mechanical datasheet.
  2. EVIDENCE PER ROW. A citation carries the page and the line it is on.
  3. THE B5 SCOPE DECISION (applicability_v2) runs on the live path over
     stored scope records - only with an owner-approved taxonomy, and a
     NOT_APPLICABLE excludes only when confirmed and never for a cited
     standard.
  4. A CITED STANDARD THAT IS NOT HELD is MISSING_LOCALLY and the run is never
     approved; AN EMPTY EVALUATION is never approved.
  5. The standards, reasons and missing ones reach the run's standards list
     and the CRS (preview and workbook).

Synthetic documents only; identifiers are public standard numbers or made up.
"""
from __future__ import annotations

import io
import json
import uuid

import pymupdf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import (access, applicability, comparison, db, keyword, standards,
                 submittal_review)
from app.config import settings
from app.crs_export import STANDARDS_SHEET
from app.main import app

NOW = "2026-09-25T00:00:00Z"


@pytest.fixture(autouse=True)
def temp_storage(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "db_path", tmp_path / "b5.sqlite")
    monkeypatch.setattr(settings, "auth_mode", access.AUTH_DISABLED)
    monkeypatch.setattr(settings, "applicability_taxonomy_path", None)
    db.reset_connection()
    db.init_db()
    submittal_review.ensure_schema()
    keyword.ensure_schema()
    yield
    app.dependency_overrides.clear()
    db.reset_connection()


def _doc(doc_id, filename, role, text="", page=1, **meta):
    with db.connect() as conn:
        conn.execute("""INSERT INTO documents
            (id,filename,sha256,size_bytes,stored_path,status,page_count,uploaded_at)
            VALUES (?,?,?,1,?,'ready',?,?)""",
            (doc_id, filename, f"sha-{doc_id}", filename, max(page, 1), NOW))
        cols = ["document_id", "suggested_by", "document_role", *meta]
        vals = [doc_id, "test", role, *meta.values()]
        conn.execute(f"INSERT INTO document_classification ({','.join(cols)})"
                     f" VALUES ({','.join('?' * len(vals))})", vals)
        conn.execute("""INSERT INTO chunks
            (id,document_id,filename,ordinal,page_start,page_end,section,kind,
             text,token_count,content_hash,retrievable)
            VALUES (?,?,?,0,?,?,NULL,'prose',?,1,?,1)""",
            (f"{doc_id}-c1", doc_id, filename, page, page, text or filename, f"h-{doc_id}"))
    keyword.index_document(doc_id)
    return doc_id


def _run(sub):
    run_id = str(uuid.uuid4())
    with db.connect() as conn:
        conn.execute("""INSERT INTO review_runs
            (id,submittal_document_id,status,created_at,updated_at)
            VALUES (?,?,'running',?,?)""", (run_id, sub, NOW, NOW))
    return run_id


def _rows(run_id):
    return {r["standard_document_id"]: dict(r) for r in db.connect().execute(
        "SELECT * FROM review_applicable_standards WHERE review_run_id = ?", (run_id,))}


def _scope(*ids):
    return frozenset(ids)


# ======================================================= 1. selection on evidence

def test_a_shared_discipline_alone_is_considered_not_applied():
    std = _doc("std_d", "mech.pdf", "COMPANY_STANDARD", discipline="Mechanical")
    sub = _doc("sub", "pump.pdf", "CONTRACTOR_SUBMITTAL", discipline="Mechanical")
    run = _run(sub)
    result = applicability.select(sub, allowed_document_ids=_scope(std, sub),
                                  review_run_id=run)
    [row] = [r for r in result["selected"] if r["standard_document_id"] == std]
    assert row["method"] == applicability.METHOD_DISCIPLINE     # it WAS considered
    assert row["included"] is False
    stored = _rows(run)[std]
    assert stored["included"] == 0
    assert "shared discipline alone" in stored["exclusion_reason"]


def test_similar_wording_alone_is_considered_not_applied():
    words = "centrifugal impeller casing mechanical seal bearing housing coupling"
    std = _doc("std_s", "similar.pdf", "COMPANY_STANDARD", text=words)
    sub = _doc("sub", "pump.pdf", "CONTRACTOR_SUBMITTAL", text=words)
    run = _run(sub)
    result = applicability.select(sub, allowed_document_ids=_scope(std, sub),
                                  review_run_id=run)
    [row] = [r for r in result["selected"] if r["standard_document_id"] == std]
    assert row["method"] == applicability.METHOD_SEMANTIC and row["included"] is False
    assert _rows(run)[std]["included"] == 0


def test_an_equipment_classification_match_is_applied():
    std = _doc("std_e", "pumps.pdf", "COMPANY_STANDARD", equipment_type="pump",
               discipline="Mechanical")
    sub = _doc("sub", "pump.pdf", "CONTRACTOR_SUBMITTAL", equipment_type="pump",
               discipline="Mechanical")
    run = _run(sub)
    applicability.select(sub, allowed_document_ids=_scope(std, sub), review_run_id=run)
    stored = _rows(run)[std]
    assert stored["included"] == 1 and stored["selection_method"] == "equipment_type"


def test_a_standard_considered_but_not_applied_is_not_compared(tmp_path):
    """The policy reaches the comparison: a discipline-only standard's
    requirement produces no finding."""
    std = _doc("std_d", "mech.pdf", "COMPANY_STANDARD", discipline="Mechanical")
    sub = _doc("sub", "pump.pdf", "CONTRACTOR_SUBMITTAL", discipline="Mechanical")
    standards.create_requirement(
        standard_document_id=std, chunk_id="std_d-c1", clause="6.1", page=1,
        requirement_text="The design pressure shall not exceed 20 barg.",
        source_text="mech.pdf",
        structured={"requirement_type": "numeric_limit", "operator": "<=", "value": 20,
                    "unit": "bar", "raw_value": "20", "raw_unit": "barg",
                    "field": "design pressure", "subject": "design pressure"})
    run = _run(sub)
    applicability.select(sub, allowed_document_ids=_scope(std, sub), review_run_id=run)
    result = comparison.run_comparison(run, allowed_document_ids=_scope(std, sub))
    assert result["requirements_evaluated"] == 0


# ================================================= 2. evidence for a citation

def test_a_citation_carries_its_page_and_line():
    std = _doc("std_610", "API-610.pdf", "COMPANY_STANDARD", document_number="API 610")
    sub = _doc("sub", "pump.pdf", "CONTRACTOR_SUBMITTAL", page=3,
               text="GENERAL NOTES\nCasing and impeller per API 610 11th edition\nEND")
    run = _run(sub)
    applicability.select(sub, allowed_document_ids=_scope(std, sub), review_run_id=run)
    stored = _rows(run)[std]
    assert stored["included"] == 1 and stored["selection_method"] == "referenced"
    assert stored["evidence_page"] == 3
    assert stored["evidence_quote"] == "Casing and impeller per API 610 11th edition"
    assert "(page 3)" in stored["selection_reason"]


def test_evidence_is_never_invented_for_a_standard_not_cited():
    std = _doc("std_e", "pumps.pdf", "COMPANY_STANDARD", equipment_type="pump")
    sub = _doc("sub", "pump.pdf", "CONTRACTOR_SUBMITTAL", equipment_type="pump",
               text="Casing per API 610")
    run = _run(sub)
    applicability.select(sub, allowed_document_ids=_scope(std, sub), review_run_id=run)
    stored = _rows(run)[std]
    assert stored["evidence_page"] is None and stored["evidence_quote"] is None


# ======================================================= 3. the scope decision

TAXONOMY = {
    "lexicon": {"centrifugal pump": ["type", "Centrifugal Pump"],
                "pump": ["family", "Pumps"],
                "rotating equipment": ["class", "Rotating"],
                "pressure vessel": ["family", "Pressure vessels"]},
    "types": {"Centrifugal Pump": {"family": "Pumps", "class": "Rotating"}},
}
COVERS = {"covered_equipment": [{"term": "centrifugal pumps", "page": 2,
                                 "quote": "This standard applies to centrifugal pumps"}],
          "explicit_exclusions": [], "explicit_limits": [], "generic_scope": False}
EXCLUDES = {"covered_equipment": [], "explicit_limits": [], "generic_scope": False,
            "explicit_exclusions": [{"term": "centrifugal pumps", "page": 1,
                                     "quote": "This standard does not apply to centrifugal pumps"}]}


@pytest.fixture
def taxonomy(tmp_path, monkeypatch):
    path = tmp_path / "taxonomy.json"
    path.write_text(json.dumps(TAXONOMY), encoding="utf-8")
    monkeypatch.setattr(settings, "applicability_taxonomy_path", path)
    return path


def _pump_submittal(**extra):
    return _doc("sub", "pump.pdf", "CONTRACTOR_SUBMITTAL",
                equipment_type="centrifugal pump", **extra)


def test_without_an_approved_taxonomy_the_scope_step_says_it_did_not_run():
    std = _doc("std_x", "x.pdf", "COMPANY_STANDARD", discipline="Mechanical")
    applicability.store_scope_record(std, COVERS)
    sub = _pump_submittal(discipline="Mechanical")
    run = _run(sub)
    result = applicability.select(sub, allowed_document_ids=_scope(std, sub),
                                  review_run_id=run)
    assert "no equipment taxonomy is approved" in result["scope_decision_not_run"]
    stored = _rows(run)[std]
    assert stored["included"] == 0                      # nothing decided it applies
    assert "no equipment taxonomy is approved" in stored["exclusion_reason"]


def test_a_scope_clause_naming_the_equipment_applies_the_standard(taxonomy):
    std = _doc("std_x", "x.pdf", "COMPANY_STANDARD")   # no v1 match at all
    applicability.store_scope_record(std, COVERS)
    sub = _pump_submittal()
    run = _run(sub)
    applicability.select(sub, allowed_document_ids=_scope(std, sub), review_run_id=run)
    stored = _rows(run)[std]
    assert stored["included"] == 1 and stored["selection_method"] == "scope"
    assert stored["scope_decision"] == "APPLICABLE"
    assert stored["evidence_page"] == 2
    assert stored["evidence_quote"] == "This standard applies to centrifugal pumps"


def test_a_confirmed_scope_exclusion_is_not_applied_and_says_why(taxonomy):
    std = _doc("std_e", "pumps.pdf", "COMPANY_STANDARD", equipment_type="centrifugal pump")
    applicability.store_scope_record(std, EXCLUDES, not_applicable_confirmed=True)
    sub = _pump_submittal()
    run = _run(sub)
    applicability.select(sub, allowed_document_ids=_scope(std, sub), review_run_id=run)
    stored = _rows(run)[std]
    assert stored["included"] == 0 and stored["scope_decision"] == "NOT_APPLICABLE"
    assert "does not apply to centrifugal pumps" in stored["exclusion_reason"]


def test_an_unconfirmed_scope_exclusion_excludes_nothing(taxonomy):
    std = _doc("std_e", "pumps.pdf", "COMPANY_STANDARD", equipment_type="centrifugal pump")
    applicability.store_scope_record(std, EXCLUDES, not_applicable_confirmed=False)
    sub = _pump_submittal()
    run = _run(sub)
    applicability.select(sub, allowed_document_ids=_scope(std, sub), review_run_id=run)
    stored = _rows(run)[std]
    assert stored["included"] == 1 and stored["scope_decision"] == "UNKNOWN"


def test_a_cited_standard_is_never_excluded_by_its_scope_reading(taxonomy):
    std = _doc("std_610", "API-610.pdf", "COMPANY_STANDARD", document_number="API 610")
    applicability.store_scope_record(std, EXCLUDES, not_applicable_confirmed=True)
    sub = _pump_submittal(text="Casing per API 610")
    run = _run(sub)
    applicability.select(sub, allowed_document_ids=_scope(std, sub), review_run_id=run)
    stored = _rows(run)[std]
    assert stored["included"] == 1 and stored["selection_method"] == "referenced"
    assert "engineer to confirm" in stored["selection_reason"]


# ================================================ 4. the code is never an approval

SUFFICIENT = {"sufficient": True, "fields_read": 40, "fields_estimated": 40}


def _f(status):
    return {"compliance_status": status}


def test_nothing_evaluated_is_never_approved():
    result = comparison.recommend_code([], SUFFICIENT)
    assert result["code"] == comparison.CODE_MANUAL
    assert "no requirement was evaluated" in result["reason"]


def test_only_not_applicable_findings_are_never_approved():
    result = comparison.recommend_code([_f(comparison.NOT_APPLICABLE)] * 3, SUFFICIENT)
    assert result["code"] == comparison.CODE_MANUAL


def test_a_cited_standard_not_held_blocks_approval():
    result = comparison.recommend_code([_f(comparison.COMPLIANT)], SUFFICIENT,
                                       missing_references=["API 682"])
    assert result["code"] == comparison.CODE_MANUAL
    assert comparison.MISSING_LOCALLY in result["reason"] and "API 682" in result["reason"]
    assert result["missing_locally"] == 1


def test_a_cited_standard_not_held_blocks_approval_with_comments_too():
    result = comparison.recommend_code([_f(comparison.MISSING_INFORMATION)], SUFFICIENT,
                                       missing_references=["API 682"])
    assert result["code"] == comparison.CODE_MANUAL


def test_a_breach_still_rejects_when_a_standard_is_also_missing():
    result = comparison.recommend_code([_f(comparison.NON_COMPLIANT)], SUFFICIENT,
                                       missing_references=["API 682"])
    assert result["code"] == comparison.CODE_REJECTED
    assert result["missing_locally"] == 1


def test_every_requirement_met_with_nothing_missing_is_still_approved():
    """The positive control: the new gates refuse only what they name."""
    result = comparison.recommend_code([_f(comparison.COMPLIANT)], SUFFICIENT)
    assert result["code"] == comparison.CODE_APPROVED


# =========================================== 5. the live route, end to end

def _pump_pdf(tmp_path, cites):
    path = tmp_path / "pump.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_text((40, 80), f"Pump shall comply with {cites}.", fontsize=7)
    page.insert_text((40, 110), "DESIGN PRESSURE", fontsize=7)
    page.insert_text((250, 110), "16 barg", fontsize=7)
    pdf.save(str(path))
    text = pdf[0].get_text()
    pdf.close()
    with db.connect() as conn:
        conn.execute("""INSERT INTO documents
            (id,filename,sha256,size_bytes,stored_path,status,page_count,uploaded_at)
            VALUES ('sub','pump.pdf','sha-sub',1,?,'ready',1,?)""", (str(path), NOW))
        conn.execute("INSERT INTO document_classification (document_id,suggested_by,"
                     "document_role) VALUES ('sub','test','CONTRACTOR_SUBMITTAL')")
        conn.execute("INSERT INTO pages (document_id,page_no,text,char_count,needs_ocr,"
                     "batch_no) VALUES ('sub',1,?,?,0,0)", (text, len(text)))
        conn.execute("""INSERT INTO chunks
            (id,document_id,filename,ordinal,page_start,page_end,section,kind,
             text,token_count,content_hash,retrievable)
            VALUES ('sub-c1','sub','pump.pdf',1,1,1,NULL,'prose',?,1,'h-sub',1)""", (text,))
    return "sub"


def _api_610_with_a_limit():
    std = _doc("std_610", "API-610.pdf", "COMPANY_STANDARD", document_number="API 610",
               text="The design pressure shall not exceed 20 barg.")
    standards.create_requirement(
        standard_document_id=std, chunk_id="std_610-c1", clause="6.3", page=1,
        requirement_text="The design pressure shall not exceed 20 barg.",
        source_text="The design pressure shall not exceed 20 barg.",
        structured={"requirement_type": "numeric_limit", "operator": "<=", "value": 20,
                    "unit": "bar", "raw_value": "20", "raw_unit": "barg",
                    "field": "design pressure", "subject": "design pressure"})
    return std


def _review(monkeypatch, tmp_path, cites, extra=frozenset()):
    # One fact on a one-page sheet: the nominal page estimate is lowered so
    # extraction completeness passes and the tests reach the gates under test.
    monkeypatch.setattr(comparison, "FIELDS_PER_PAGE_NOMINAL", 1)
    with db.connect() as conn:
        conn.execute("INSERT OR IGNORE INTO users (id,email,display_name,password_hash,"
                     "is_active,created_at) VALUES ('eng','eng@e.test','Eng','x',1,?)", (NOW,))
    std = _api_610_with_a_limit()
    sub = _pump_pdf(tmp_path, cites)
    app.dependency_overrides[access.current_scope] = lambda: access.AccessScope(
        user_id="eng", allowed_document_ids=frozenset({std, sub, *extra}))
    client = TestClient(app)
    response = client.post("/api/reviews/run", json={"submittal_document_id": sub})
    assert response.status_code == 200, response.text
    run_id = response.json()["review_run_id"]
    # P3: the route queues; the worker runs it. Drained here exactly as the
    # worker would.
    from app import job_queue, review_jobs
    job_id = review_jobs.claim_next(job_queue.worker_id())
    assert review_jobs.run(job_id, job_queue.worker_id()) == "done"
    return client, run_id, std


def _held(doc_id, number):
    """A held standard with no requirement: cited, applied, nothing to compare."""
    return _doc(doc_id, f"{doc_id}.pdf", "COMPANY_STANDARD", document_number=number)


def test_a_missing_standard_is_named_when_it_is_why_completeness_is_low(monkeypatch, tmp_path):
    """1 of 2 cited standards held: completeness 0.5 gates the code, and the
    reason names the missing standard as a cause."""
    _client, run_id, std = _review(monkeypatch, tmp_path, "API 610 and API 682")
    outcome = comparison.run_outcome(run_id, allowed_document_ids=frozenset({std, "sub"}))
    assert outcome["recommended_code"] == comparison.CODE_MANUAL
    assert comparison.MISSING_LOCALLY in outcome["reason"] and "API 682" in outcome["reason"]


def test_the_live_route_never_approves_with_a_cited_standard_missing(monkeypatch, tmp_path):
    """3 of 4 cited standards held: completeness passes (0.75), the one
    compared requirement is met - and the missing fourth still stops the
    approval. This is the MISSING_LOCALLY gate itself, not completeness."""
    held = {_held("std_b73", "ASME B73.1"), _held("std_nace", "NACE MR0175")}
    client, run_id, std = _review(monkeypatch, tmp_path,
                                  "API 610, ASME B73.1, NACE MR0175 and API 682",
                                  extra=held)
    scope = frozenset({std, "sub", *held})
    outcome = comparison.run_outcome(run_id, allowed_document_ids=scope)
    findings = submittal_review.list_run_findings(run_id, allowed_document_ids=scope)
    assert [f["compliance_status"] for f in findings] == [comparison.COMPLIANT]
    assert outcome["completeness"]["sufficient"] is True
    assert outcome["recommended_code"] == comparison.CODE_MANUAL
    assert outcome["reason"].startswith("Manual review: 1 standard(s) the submittal cites")
    assert "API 682" in outcome["reason"]
    assert outcome["missing_references"] == [
        {"identifier": "API 682", "status": comparison.MISSING_LOCALLY}]


def test_the_live_route_approves_when_nothing_is_missing(monkeypatch, tmp_path):
    """The positive control on the same route: without the missing citation
    the same compliant sheet is approved, so the refusal above is the gate."""
    _client, run_id, std = _review(monkeypatch, tmp_path, "API 610")
    outcome = comparison.run_outcome(run_id, allowed_document_ids=frozenset({std, "sub"}))
    assert outcome["recommended_code"] == comparison.CODE_APPROVED


def test_the_standards_and_reasons_reach_the_run_and_the_crs(monkeypatch, tmp_path):
    client, run_id, std = _review(monkeypatch, tmp_path, "API 610 and API 682")

    listed = client.get(f"/api/reviews/runs/{run_id}/standards").json()
    [row] = listed["standards"]
    assert row["standard_document_id"] == std and row["included"] is True
    assert row["selection_method"] == "referenced" and row["evidence_page"] == 1
    assert "API 610" in row["evidence_quote"]
    assert listed["missing_references"] == [
        {"identifier": "API 682", "status": comparison.MISSING_LOCALLY}]

    preview = client.get(f"/api/reviews/runs/{run_id}/crs/preview").json()
    by_name = {s["standard"]: s for s in preview["applicable_standards"]}
    assert by_name["API-610.pdf"]["status"] == "Applied"
    assert "named in the submittal as API 610" in by_name["API-610.pdf"]["reason"]
    assert by_name["API 682"]["status"] == comparison.MISSING_LOCALLY

    workbook = load_workbook(io.BytesIO(client.get(f"/api/reviews/runs/{run_id}/crs").content))
    sheet = workbook[STANDARDS_SHEET]
    printed = [[c.value for c in r] for r in sheet.iter_rows(min_row=2)]
    assert ["API-610.pdf", "Applied"] == printed[0][:2]
    assert ["API 682", comparison.MISSING_LOCALLY] == printed[1][:2]
    assert workbook.worksheets[0].title == "CRS"        # the client's sheet stays first


def test_the_applicability_list_reports_a_discipline_match_as_unknown_not_applicable():
    """The four-bucket list (`applicability_with_reasons`) reads the same
    policy: a considered-not-applied standard is not "applicable"."""
    std = _doc("std_d", "mech.pdf", "COMPANY_STANDARD", discipline="Mechanical")
    sub = _doc("sub", "pump.pdf", "CONTRACTOR_SUBMITTAL", discipline="Mechanical")
    [entry] = applicability.applicability_with_reasons(
        sub, allowed_document_ids=_scope(std, sub))
    assert entry["status"] == applicability.STATUS_UNKNOWN
    assert "shared discipline alone" in entry["reason"]


def test_an_empty_taxonomy_setting_is_no_taxonomy(monkeypatch, tmp_path):
    """`APPLICABILITY_TAXONOMY_PATH=` in .env parses as Path("."), a directory."""
    from pathlib import Path
    monkeypatch.setattr(settings, "applicability_taxonomy_path", Path("."))
    assert applicability.load_taxonomy() is None
