"""#747: a DRAFT answer key from a real Comment Resolution Sheet.

The real CRSs KJO engineers wrote are the answer key. Builders never read
them; this converter is built and tested on INVENTED sheets rendered in the
client's own template (`crs_export.build_crs`). Each comment becomes one
defect item with the standard and clause it names; a comment whose standard
or clause cannot be read is never guessed, it is listed by row for the merger
to complete. Invented comments only.
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import crs_export
from tools import review_score as rs

REPO = Path(__file__).resolve().parents[2]


def _sheet(comments: list[tuple[str, str]]) -> list[tuple]:
    findings = [{"document_name": "invented-psv.pdf", "page_section": page, "comment": text,
                 "comment_by": "Engineer A"} for page, text in comments]
    data = crs_export.build_crs(findings, {"review_run_id": "invented"})
    return list(load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
                .iter_rows(values_only=True))


COMMENTS = [
    ("p.2 - Set pressure", "Set pressure exceeds the limit of API 520-I 5.3.3; please revise."),
    ("p.3 - Margin", "Per SAES-J-600 clause 8, the margin above operating pressure is too small."),
    ("p.4 - Orifice area", "Provide the orifice area (API RP 520 Part I, section 5.2)."),
    ("p.5 - Notes", "Please clarify note 4 on the drawing."),            # no standard
    ("p.6 - Bonnet", "Bonnet type to follow API 526."),                  # standard, no clause
]


def test_each_comment_with_a_standard_and_a_clause_becomes_a_defect_item():
    key, todo = rs.key_from_crs(_sheet(COMMENTS), name="invented-psv", approved_by="Engineer A")
    rs.validate_key(key)                                                 # a valid key for the scorer
    got = {(i["standard"], i["clause"], i.get("field")) for i in key["items"]}
    assert got == {("API 520-I", "5.3.3", "Set pressure"), ("SAES-J-600", "8", "Margin"),
                   ("API RP 520 Part I", "5.2", "Orifice area")}
    assert all(i["kind"] == "defect" and "drafted by code" in i["note"] for i in key["items"])
    assert key["source"] == "engineer_confirmed" and key["approved_by"] == "Engineer A"


def test_a_comment_whose_standard_or_clause_cannot_be_read_is_listed_not_guessed():
    _key, todo = rs.key_from_crs(_sheet(COMMENTS), name="invented-psv", approved_by="Engineer A")
    assert sorted(t["missing"] for t in todo) == ["clause", "standard"]
    assert all(set(t) == {"row", "item_no", "missing"} for t in todo)    # rows only, never the comment text


def test_a_sheet_that_is_not_a_crs_or_has_no_approver_is_refused():
    with pytest.raises(rs.KeyError_, match="COMPANY Comments"):
        rs.key_from_crs([("a", "b"), ("c", "d")], name="x", approved_by="Engineer A")
    with pytest.raises(rs.KeyError_, match="approved_by"):
        rs.key_from_crs(_sheet(COMMENTS), name="x", approved_by=" ")


def test_the_key_names_the_standard_as_written_and_still_matches_the_librarys_file():
    """The scorer matched standards by substring only: 'API RP 520 Part I' never
    matched a finding from 'API-520-I.pdf'. The one identifier rule now decides."""
    item = {"standard": "API RP 520 Part I", "clause": "5.2"}
    assert rs._standard_matches({"standard": "API-520-I.pdf"}, item)
    assert not rs._standard_matches({"standard": "API-520-II.pdf"}, item)
    assert rs._standard_matches({"standard": "x", "standard_number": "SAES-J-600"},
                                {"standard": "SAES-J-600", "clause": "8"})


def test_the_command_refuses_to_write_a_drafted_key_inside_the_repo(tmp_path):
    sys.path.insert(0, str(REPO))
    from scripts import review_score as cli
    crs = tmp_path / "crs.xlsx"
    crs.write_bytes(crs_export.build_crs(
        [{"document_name": "d.pdf", "page_section": "p.2 - Set pressure", "comment": COMMENTS[0][1],
          "comment_by": "Engineer A"}], {}))
    assert cli.main(["from-crs", "--crs", str(crs), "--name", "x", "--approved-by", "A",
                     "--out", str(REPO / "eval" / "review" / "keys" / "drafted.json")]) == 2
    assert not (REPO / "eval" / "review" / "keys" / "drafted.json").exists()
    out = tmp_path / "drafted.json"
    assert cli.main(["from-crs", "--crs", str(crs), "--name", "x", "--approved-by", "A",
                     "--out", str(out)]) == 0
    assert json.loads(out.read_text(encoding="utf-8"))["items"][0]["clause"] == "5.3.3"
