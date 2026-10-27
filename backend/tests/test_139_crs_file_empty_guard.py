"""Order 139 part 2, item 5 (spec sections e and f) in the FILE: the Priority
column (it replaced Severity, which said "minor" on 324 of 329 rows), the
empty-comment guard on both copies, and the issue gate's blocker for a
confirmed row with no text. Invented data only.
"""

from __future__ import annotations

import io

import openpyxl

from app import crs_export as ce
from app import crs_issue
from app import crs_mapping as m
from app.schemas import CrsPreview

COVERED = {"pages_total": 1, "pages_not_read_into_fields": []}


def _row(n: int, **over) -> dict:
    return {
        "finding_id": f"f{n}",
        "document_name": "inv.pdf",
        "page_section": f"p.{n}",
        "comment": f"Invented comment {n}.",
        "comment_by": "AI Review",
        "row_kind": "missing_information",
        "severity": "minor",
        "engineer_confirmed": True,
        "priority": m.PRIORITY_MISSING_VALUE,
        **over,
    }


def _sheet(rows: list[dict], copy: str):
    book = openpyxl.load_workbook(io.BytesIO(ce.build_crs(rows, {"copy": copy})))
    return book["CRS"]


# ------------------------------------------------------------ f: Priority


def test_priority_column_prints_the_rank_word():
    """Column I is "Priority" in both copies and prints the row's rank word,
    never its severity ("minor" said nothing)."""
    rows = [
        _row(1, priority=m.PRIORITY_OUT_OF_LIMIT, row_kind="non_compliant"),
        _row(2),
    ]
    for copy in ce.COPIES:
        view = ce.build_crs_view(rows, {"copy": copy})
        assert view["columns"][8] == ce.PRIORITY_COLUMN == "Priority"
        assert "Severity" not in view["columns"]
        assert [r["priority_label"] for r in view["rows"]] == [
            "Out of limit",
            "Missing value",
        ]
        assert [r["severity"] for r in view["rows"]] == [
            "minor",
            "minor",
        ]  # kept, not printed
        ws = _sheet(rows, copy)
        assert ws.cell(row=ce.COLUMN_HEADER_ROW, column=9).value == "Priority"
        assert [ws.cell(row=ce.FIRST_DATA_ROW + i, column=9).value for i in (0, 1)] == [
            "Out of limit",
            "Missing value",
        ]


def test_issue_copy_says_clarification_for_rank_three():
    """To the contractor an engineer-to-check row is a request to clarify,
    not an internal to-do; the internal copy keeps the engineer's words."""
    rows = [
        _row(1, priority=m.PRIORITY_ENGINEER_TO_CHECK, row_kind="needs_engineer_review")
    ]
    internal = ce.build_crs_view(rows, {"copy": ce.COPY_INTERNAL})["rows"][0]
    issue = ce.build_crs_view(rows, {"copy": ce.COPY_ISSUE})["rows"][0]
    assert internal["priority_label"] == "Engineer to check"
    assert issue["priority_label"] == "Clarification"
    assert (
        _sheet(rows, ce.COPY_ISSUE).cell(row=ce.FIRST_DATA_ROW, column=9).value
        == "Clarification"
    )


def test_row_without_priority_prints_nothing():
    """A row from a caller that sent no rank prints an empty Priority cell -
    never a guessed word. A carried-forward row's kind IS its rank (main
    appends it from an earlier comment's snapshot, which holds no kind)."""
    rows = [
        _row(1, priority=None),
        _row(2, priority=None, row_kind=m.ROW_KIND_CARRIED_FORWARD),
    ]
    view = ce.build_crs_view(rows, {"copy": ce.COPY_INTERNAL})
    assert [(r["priority"], r["priority_label"]) for r in view["rows"]] == [
        (None, ""),
        (m.PRIORITY_CARRIED_FORWARD, "Carried forward"),
    ]
    assert _sheet(rows, ce.COPY_INTERNAL).cell(
        row=ce.FIRST_DATA_ROW, column=9
    ).value in (None, "")


def test_file_keeps_the_sheets_order():
    """The view never re-sorts: the order `build_crs_rows` decided (rank
    first) is the order of the file and the preview."""
    rows = [_row(3, comment="c"), _row(1, comment="a"), _row(2, comment="b")]
    view = ce.build_crs_view(rows, {"copy": ce.COPY_INTERNAL})
    assert [r["page_section"] for r in view["rows"]] == ["p.3", "p.1", "p.2"]


# ------------------------------------------- e: no empty COMPANY Comments


def test_issue_copy_never_carries_a_row_with_no_text():
    """A confirmed row with no comment text is never issued - and the copy
    counts it (`rows_withheld_no_text`), never silently."""
    rows = [_row(1), _row(2, comment=""), _row(3, comment="   ")]
    view = ce.build_crs_view(rows, {"copy": ce.COPY_ISSUE})
    assert [r["page_section"] for r in view["rows"]] == ["p.1"]
    assert [r["item_no"] for r in view["rows"]] == [1]  # numbering stays 1..N
    assert view["rows_withheld_no_text"] == 2
    ws = _sheet(rows, ce.COPY_ISSUE)
    assert ws.cell(row=ce.FIRST_DATA_ROW + 1, column=4).value in (
        None,
        "",
    )  # no second row
    # the preview route serves the view through this model: the count survives
    assert CrsPreview.model_validate(view).rows_withheld_no_text == 2
    assert (
        ce.build_crs_view([_row(1)], {"copy": ce.COPY_ISSUE})["rows_withheld_no_text"]
        == 0
    )


def test_internal_copy_names_a_row_with_no_text():
    """The internal copy keeps the row and says why its cell has no comment,
    under the row's reference; the row is marked `comment_missing`."""
    rows = [_row(1), _row(2, comment="")]
    view = ce.build_crs_view(rows, {"copy": ce.COPY_INTERNAL})
    assert view["rows_withheld_no_text"] == 0
    assert [r["comment_missing"] for r in view["rows"]] == [False, True]
    assert view["rows"][1]["comment"] == ce.EMPTY_COMMENT_TEXT  # order 236: no Ref line
    cell = (
        _sheet(rows, ce.COPY_INTERNAL).cell(row=ce.FIRST_DATA_ROW + 1, column=4).value
    )
    assert cell.endswith(ce.EMPTY_COMMENT_TEXT)
    numbered = ce.build_crs_view(
        [_row(2, comment="", crs_ref="CRS-INV-001")], {"copy": "internal"}
    )
    # 1b: a NUMBERED row was issued - it says which comment and what to do
    assert numbered["rows"][0]["comment"] == ce.empty_comment_text("CRS-INV-001")
    assert numbered["rows"][0]["comment"].startswith("Comment CRS-INV-001 was issued")
    assert "An engineer must write it" not in numbered["rows"][0]["comment"]


def test_an_unconfirmed_ai_row_is_internal_with_its_waiting_sentence():
    """End to end from the mapping: the AI item's COMPANY Comments is the
    waiting sentence (not EMPTY_COMMENT_TEXT, not the draft), and the issue
    copy still leaves the unconfirmed row out."""
    ai = {
        "id": "ai",
        "origin": "ai_engineering_check",
        "compliance_status": "NEEDS_ENGINEER_REVIEW",
        "finding": "Confirm the invented guard is fitted.",
        "contractor_page": 1,
    }
    rows = m.build_crs_rows([ai], [], "inv.pdf")
    internal = ce.build_crs_view(rows, {"copy": ce.COPY_INTERNAL})["rows"][0]
    assert internal["comment"].endswith(m.AI_UNCONFIRMED_COMMENT)
    assert internal["comment_missing"] is False
    assert internal["priority_label"] == "Engineer to check"
    issue = ce.build_crs_view(rows, {"copy": ce.COPY_ISSUE})
    assert issue["rows"] == [] and issue["rows_withheld_no_text"] == 0


# -------------------------------------------------------- e: the issue gate


def test_gate_blocks_a_confirmed_row_with_no_text():
    """The export's drop is a last guard; the gate is what normally stops a
    confirmed row with an empty comment from being issued."""
    rows = [
        {
            "finding_id": "f1",
            "row_kind": "missing_information",
            "engineer_confirmed": True,
            "comment": "",
        },
        {
            "finding_id": "f2",
            "row_kind": "missing_information",
            "engineer_confirmed": True,
            "comment": "Contractor to provide the invented value.",
        },
    ]
    gate = crs_issue.open_items("run-1", rows, COVERED)
    assert gate["issuable"] is False
    assert gate["blockers"] == [crs_issue.no_text_blocker(["on row 1"])]
    assert gate["blockers"][0].startswith("Comment on row 1 has no text")


def test_gate_ignores_rows_the_issue_copy_would_not_carry():
    """Only a row that would reach the contractor counts: not an unconfirmed
    draft (that is an undecided finding already), not a safety tick,
    not one an engineer rejected. (A confirmed row that still holds an AI
    draft is NOT ignored since 1b: see test_1b_issued_ai_comment.)"""
    rows = [
        {
            "finding_id": "t",
            "row_kind": m.ROW_KIND_SAFETY_TICK,
            "engineer_confirmed": True,
            "comment": "",
        },
        {
            "finding_id": "r",
            "row_kind": "standard_not_cited",
            "engineer_confirmed": True,
            "engineer_rejected": True,
            "comment": "",
        },
    ]
    gate = crs_issue.open_items("run-1", rows, COVERED)
    assert gate["blockers"] == [] and gate["issuable"] is True


# ------------------------------- #1178 review: the waiting sentence alone


def test_a_carried_forward_waiting_sentence_is_never_issued():
    """A waiting sentence is text only beside its draft. Stored under an
    issued number and carried forward with no AI Review Comments beside it,
    it would tell the contractor something false: the issue copy leaves it
    out and counts it, the internal copy says no text was produced, and the
    gate blocks the issue."""
    for sentence in (m.AI_UNCONFIRMED_COMMENT, m.WEB_UNCONFIRMED_COMMENT):
        carried = _row(
            9, comment=sentence, row_kind=m.ROW_KIND_CARRIED_FORWARD, priority=None
        )
        issue = ce.build_crs_view([_row(1), carried], {"copy": ce.COPY_ISSUE})
        assert [r["page_section"] for r in issue["rows"]] == ["p.1"]
        assert issue["rows_withheld_no_text"] == 1
        internal = ce.build_crs_view([carried], {"copy": ce.COPY_INTERNAL})["rows"][0]
        assert internal["comment_missing"] is True
        assert sentence not in internal["comment"]
        gate = crs_issue.open_items("run-1", [carried], COVERED)
        # a carried-forward comment: closing it also clears it (1b)
        assert gate["blockers"] == [
            crs_issue.no_text_blocker(["on row 1"], ["on row 1"])
        ]
    # beside its draft the sentence IS the row's text (internal copy)
    draft = _row(2, comment=m.AI_UNCONFIRMED_COMMENT, ai_review_comment="a draft")
    assert m.row_has_no_text(draft) is False
