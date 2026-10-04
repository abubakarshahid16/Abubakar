"""r2 S2 (CLAUDE.md rule 5): a grant on a SUBMITTAL is not a grant on the
STANDARD it was compared with.

Findings carry the standard's id, clause, page, requirement wording and
filename. A caller who may read only the submittal must see the verdict and
none of that - on the findings list, the finding's traceability, the CRS
preview, the CRS workbook, and the reply to a PATCH. A caller granted both
sees it all, unchanged. Mutations: scripts/mutations/r2_security.py M1981-M1982.
"""

from __future__ import annotations

import io
import json
import uuid

from openpyxl import load_workbook

from app.db import connect
from app.review import STANDARD_WITHHELD

from tests.r2_security_support import h, now, temp_storage, world  # noqa: F401

#: Every string that exists only because the standard was read.
SECRETS = ("6,900 kPa", "STD-A-001", "6.2.2", "design pressure shall")


def _run_with_finding(finding_id: str = "f1") -> str:
    run_id = str(uuid.uuid4())
    with connect() as conn:
        conn.execute(
            "INSERT INTO review_runs (id, submittal_document_id, status, refusal_reason,"
            " created_at, updated_at) VALUES (?, 'doc_sub', 'completed', ?, ?, ?)",
            (run_id, json.dumps({"recommended_code": "Manual Review Required",
                                 "reason": "examined"}), now(), now()))
        row = {
            "id": finding_id, "document_id": "doc_sub", "review_run_id": run_id,
            "category": "requirement_deviation", "severity": "major",
            "requirement": "The design pressure shall be 6,900 kPa (STD-A-001 cl. 6.2.2)",
            "finding": "Outside the STD-A-001 limit of 6,900 kPa",
            "required_action": "Revise and resubmit",
            "compliance_status": "NON_COMPLIANT",
            "requirement_source_text": "The design pressure shall be 6,900 kPa.",
            "contractor_evidence_text": "2.2 bar (ga)",
            "ai_rationale": "unit_mismatch against STD-A-001 6.2.2",
            "standard_document_id": "doc_std", "standard_clause": "6.2.2",
            "standard_page": 14, "contractor_page": 4,
            "governing_sources": json.dumps(["STD-A-001"]),
            "citation_ids": json.dumps(["chunk-of-the-standard"]),
            "unresolved_evidence": "[]", "status": "open",
            "approval_status": "pending", "escalation_level": 0,
            "created_at": now(), "updated_at": now(),
        }
        conn.execute(f"INSERT INTO review_findings ({','.join(row)})"
                     f" VALUES ({','.join('?' * len(row))})", list(row.values()))
    return run_id


def _no_secret(text: str, where: str) -> None:
    for secret in SECRETS:
        assert secret.lower() not in text.lower(), f"{secret!r} leaked through {where}"


def test_findings_list_withholds_the_standard_from_a_submittal_only_reader(world):
    _run_with_finding()
    r = world.get("/api/reviews/findings", headers=h("u_sub"))
    assert r.status_code == 200
    [finding] = r.json()["findings"]
    _no_secret(json.dumps(finding), "GET /api/reviews/findings")
    assert finding["standard_document_id"] is None
    assert finding["standard_clause"] is None and finding["standard_page"] is None
    assert finding["requirement"] == STANDARD_WITHHELD
    assert finding["standard_withheld"] is True
    # The verdict, and the contractor's own side, are still there.
    assert finding["compliance_status"] == "NON_COMPLIANT"
    assert finding["severity"] == "major"
    assert finding["contractor_evidence_text"] == "2.2 bar (ga)"


def test_the_full_grant_and_the_admin_see_the_standard_unchanged(world):
    _run_with_finding()
    for user in ("u_full", "u_admin"):
        [finding] = world.get("/api/reviews/findings", headers=h(user)).json()["findings"]
        assert finding["standard_document_id"] == "doc_std", user
        assert finding["standard_clause"] == "6.2.2" and finding["standard_page"] == 14
        assert "6,900 kPa" in finding["requirement"]
        assert finding["standard_withheld"] is False


def test_traceability_withholds_the_standard(world):
    _run_with_finding()
    r = world.get("/api/reviews/findings/f1/traceability", headers=h("u_sub"))
    assert r.status_code == 200
    _no_secret(r.text, "traceability")
    assert r.json()["finding"]["compliance_status"] == "NON_COMPLIANT"
    full = world.get("/api/reviews/findings/f1/traceability", headers=h("u_full"))
    assert full.json()["finding"]["standard_clause"] == "6.2.2"


def test_the_patch_reply_withholds_the_standard(world):
    _run_with_finding()
    r = world.patch("/api/reviews/findings/f1", json={"status": "in_progress"},
                    headers=h("u_sub"))
    assert r.status_code == 200, r.text
    _no_secret(r.text, "PATCH /api/reviews/findings/{id}")


def test_the_crs_preview_withholds_the_standard(world):
    run_id = _run_with_finding()
    r = world.get(f"/api/reviews/runs/{run_id}/crs/preview", headers=h("u_sub"))
    assert r.status_code == 200, r.text
    _no_secret(r.text, "crs/preview")
    # And the preview still has the row: only the standard's words are gone.
    assert r.json()["rows"], "the preview lost the finding altogether"
    full = world.get(f"/api/reviews/runs/{run_id}/crs/preview", headers=h("u_full"))
    assert "STD-A-001" in full.text, "the full grant should still see the standard"


def test_the_crs_workbook_withholds_the_standard(world):
    run_id = _run_with_finding()
    r = world.get(f"/api/reviews/runs/{run_id}/crs", headers=h("u_sub"))
    assert r.status_code == 200, r.text
    sheet_text = " ".join(
        str(c.value) for ws in load_workbook(io.BytesIO(r.content)).worksheets
        for row in ws.iter_rows() for c in row if c.value is not None)
    _no_secret(sheet_text, "the CRS .xlsx")
    full = world.get(f"/api/reviews/runs/{run_id}/crs", headers=h("u_full"))
    full_text = " ".join(
        str(c.value) for ws in load_workbook(io.BytesIO(full.content)).worksheets
        for row in ws.iter_rows() for c in row if c.value is not None)
    assert "STD-A-001" in full_text
