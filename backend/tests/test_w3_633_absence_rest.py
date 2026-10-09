"""#633: the remaining absence paths. A sheet for a run that could not check
everything never looks complete; a failed optional check is stored as "could
not be checked"; findings written before a failure are marked partial; a run
with large unchecked parts cannot be approved; silent skips in the datasheet
self-checks, PDF export and analysis are said. Synthetic data only.
Mutations M3401-M3420."""
from __future__ import annotations

import json
import secrets

import pytest
from fastapi.testclient import TestClient

from app import (absence, access, ai_engineering_check, analysis, auth, comparison, crs_export,
                 datasheet_checks, datasheets, db, job_queue, keyword, main as main_mod,
                 review, review_jobs, standards, submittal_review, web_standards)
from app.config import settings

NOW = "2026-09-26T00:00:00Z"
W = "w-test"


@pytest.fixture(autouse=True)
def temp_storage(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "db_path", tmp_path / "w633.sqlite")
    monkeypatch.setattr(settings, "auth_mode", access.AUTH_REQUIRED)
    monkeypatch.setattr(settings, "auth_secret", secrets.token_urlsafe(48))
    monkeypatch.setattr(settings, "geometry_reader_enabled", False)
    monkeypatch.setattr(settings, "job_max_running", 1)
    db.reset_connection()
    db.init_db()
    submittal_review.ensure_schema()
    submittal_review.migrate_facts_to_per_document()
    keyword.ensure_schema()
    access.set_user_resolver(auth.resolve_user_id)
    yield
    access.set_user_resolver(None)
    monkeypatch.setattr(settings, "auth_mode", access.AUTH_DISABLED)
    db.reset_connection()


def _doc(doc_id, role, text):
    with db.connect() as conn:
        conn.execute("INSERT INTO documents (id,filename,sha256,size_bytes,stored_path,status,"
                     "page_count,uploaded_at) VALUES (?,?,?,1,'/x','ready',1,?)",
                     (doc_id, f"{doc_id}.pdf", f"sha-{doc_id}", NOW))
        conn.execute("INSERT INTO document_classification (document_id,suggested_by,document_role)"
                     " VALUES (?,?,?)", (doc_id, "test", role))
        conn.execute("INSERT INTO chunks (id,document_id,filename,ordinal,page_start,page_end,section,"
                     "kind,text,token_count,content_hash,retrievable) VALUES (?,?,?,0,1,1,'1',"
                     "'prose',?,1,?,1)", (f"{doc_id}-c1", doc_id, f"{doc_id}.pdf", text, f"h-{doc_id}"))
    keyword.index_document(doc_id)
    return doc_id


def _user(user_id, docs, admin=False):
    with db.connect() as conn:
        conn.execute("INSERT INTO users (id,email,display_name,password_hash,is_active,created_at)"
                     " VALUES (?,?,?,'x',1,?)", (user_id, f"{user_id}@x", user_id, NOW))
        rid = f"r-{user_id}"
        conn.execute("INSERT INTO roles (id,name,description,kind,created_at) VALUES (?,?,'',?,?)",
                     (rid, rid, "discipline", NOW))
        conn.execute("INSERT INTO user_roles (user_id,role_id,granted_at) VALUES (?,?,?)", (user_id, rid, NOW))
        if admin:
            conn.execute("INSERT OR IGNORE INTO roles (id,name,description,kind,created_at)"
                         " VALUES ('role_admin','admin','','capability',?)", (NOW,))
            adm = conn.execute("SELECT id FROM roles WHERE name = 'admin'").fetchone()["id"]
            conn.execute("INSERT INTO user_roles (user_id,role_id,granted_at) VALUES (?,?,?)",
                         (user_id, adm, NOW))
        for d in docs:
            conn.execute("INSERT INTO document_role_access (document_id,role_id,granted_at)"
                         " VALUES (?,?,?)", (d, rid, NOW))
    return {"Authorization": f"Bearer {auth.issue_token(user_id)}"}


@pytest.fixture
def world():
    std = _doc("std_api", "COMPANY_STANDARD", "Equipment noise level shall not exceed 90 dB(A).")
    sub = _doc("sub_pump", "CONTRACTOR_SUBMITTAL", "Pump shall comply with API 610. Noise level 95 dB(A)")
    with db.connect() as conn:
        conn.execute("UPDATE document_classification SET document_number = 'API 610' WHERE document_id = ?", (std,))
    standards.extract_requirements(std, allowed_document_ids=frozenset({std}))
    datasheets.create_fact(submittal_document_id=sub, chunk_id=f"{sub}-c1",
                           field_label="Noise level", raw_value="95", unit="dB(A)", page=1)
    admin = _user("adm", [std, sub], admin=True)
    return {"std": std, "sub": sub, "admin": admin, "scope": frozenset({std, sub})}


def _enqueue(w, user="adm"):
    return review_jobs.enqueue(w["sub"], allowed_document_ids=w["scope"], requested_by=user)


def _run_row(run_id):
    return dict(db.connect().execute("SELECT * FROM review_runs WHERE id = ?", (run_id,)).fetchone())


def _job(job_id):
    return dict(db.connect().execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone())


def _findings(run_id):
    return db.connect().execute("SELECT COUNT(*) FROM review_findings WHERE review_run_id = ?",
                                (run_id,)).fetchone()[0]


def _audits(action):
    return [dict(r) for r in db.connect().execute("SELECT * FROM audit_events WHERE action = ?", (action,))]


def _run_row(run_id):
    return dict(db.connect().execute("SELECT * FROM review_runs WHERE id = ?", (run_id,)).fetchone())


def _scope_of(w):
    return access.AccessScope(user_id="adm", allowed_document_ids=w["scope"])


@pytest.fixture
def world():
    std = _doc("std_api", "COMPANY_STANDARD", "Equipment noise level shall not exceed 90 dB(A).")
    sub = _doc("sub_pump", "CONTRACTOR_SUBMITTAL", "Pump shall comply with API 610. Noise level 95 dB(A)")
    with db.connect() as conn:
        conn.execute("UPDATE document_classification SET document_number = 'API 610' WHERE document_id = ?", (std,))
    standards.extract_requirements(std, allowed_document_ids=frozenset({std}))
    datasheets.create_fact(submittal_document_id=sub, chunk_id=f"{sub}-c1",
                           field_label="Noise level", raw_value="95", unit="dB(A)", page=1)
    admin = _user("adm", [std, sub], admin=True)
    return {"std": std, "sub": sub, "admin": admin, "scope": frozenset({std, sub})}


def _enqueue(w):
    return review_jobs.enqueue(w["sub"], allowed_document_ids=w["scope"], requested_by="adm")


# ------------------------------------------------ the one list of unchecked parts

def test_unchecked_parts_names_every_part_and_a_clean_run_has_none():
    outcome = {
        "error": "the model host went away", "partial_findings": 3,
        "standards_not_checked": ["SAES-X.pdf"],
        "table_values_not_compared": [{"count": 40}, {"count": 2}],
        "requirements_not_applied": [{"count": 7}],
        "requirements_held_back": {"definition": 1, "text_quality": 4},
        "page_coverage": {"pages_not_read_into_fields": [4, 5]},
    }
    parts = absence.unchecked_parts(
        run_status="failed", outcome=outcome, partial_findings=3,
        ai_status=absence.check_failed_status("AI engineering check", "no key"),
        web_status=absence.check_failed_status("Web standards check", "offline"))
    lines = " | ".join(p["line"] for p in parts)
    assert {p["part"] for p in parts} == {
        "run_not_completed", "partial_findings", "standards_not_checked",
        "table_values_not_compared", "text_quality_held_back",
        "unread_pages", "ai_check", "web_check", "unchecked_share"}
    assert "did not complete (status: failed): the model host went away" in lines
    assert "42 standards-table value(s) were not compared" in lines
    # #678: requirements about other equipment DO NOT APPLY; they are not an
    # unchecked part and no longer make a sheet incomplete.
    assert "were not applied" not in lines
    assert "4 requirement(s) were held back" in lines
    # The control: a completed run that left nothing unchecked has no parts.
    assert absence.unchecked_parts(run_status="completed", outcome={}) == []
    assert absence.notice_for([]) == ""
    assert absence.notice_for(parts).startswith("REVIEW INCOMPLETE: 9 part(s)")


def test_the_workbook_prints_the_incomplete_notice_above_the_code():
    import io
    import openpyxl
    meta = {"incomplete_notice": absence.notice_for([{"part": "x", "line": "y"}]),
            "recommended_code": "Manual Review Required", "unchecked_parts": ["y"]}
    cells = [c.value for row in openpyxl.load_workbook(io.BytesIO(crs_export.build_crs([], meta))).active.iter_rows()
             for c in row if c.value]
    assert any(str(v).startswith("REVIEW INCOMPLETE") for v in cells)
    clean = [c.value for row in openpyxl.load_workbook(io.BytesIO(crs_export.build_crs([], {})))
             .active.iter_rows() for c in row if c.value]
    assert not any("INCOMPLETE" in str(v) for v in clean)


# ----------------------------------------- a failed run: partial, and the sheet says so

def test_findings_written_before_a_failure_are_marked_partial_and_the_crs_says_incomplete(world, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("synthetic failure after the findings were written")
    monkeypatch.setattr(datasheet_checks, "evaluate", boom)
    run_id, job_id = _enqueue(world)
    review_jobs.claim_next(W)
    assert review_jobs.run(job_id, W) == job_queue.RETRYING
    monkeypatch.setattr(settings, "job_max_retries", 0)
    with db.connect() as conn:
        conn.execute("UPDATE jobs SET state = 'running', claimed_by = ? WHERE id = ?", (W, job_id))
    assert review_jobs.run(job_id, W) == job_queue.POISONED

    run = _run_row(run_id)
    stored = json.loads(run["refusal_reason"])
    assert run["status"] == "failed"
    assert stored["partial"] is True and stored["partial_findings"] >= 1     # positive first

    rows, meta, _name, _stamp = main_mod._crs_content(run_id, _scope_of(world), "internal")
    assert meta["incomplete_notice"].startswith("REVIEW INCOMPLETE")
    assert any("did not complete (status: failed)" in line for line in meta["unchecked_parts"])
    assert any("were written before the review stopped" in line for line in meta["unchecked_parts"])
    assert any(n["note"] == "Not checked" for n in meta["review_notes"])
    # The contractor's copy carries the notice too, never the detail.
    _r, issue_meta, _n, _s = main_mod._crs_content(run_id, _scope_of(world), "issue")
    view = crs_export.build_crs_view(rows, {**issue_meta, "copy": "issue"})
    assert view["incomplete_notice"].startswith("REVIEW INCOMPLETE") and view["unchecked_parts"] == []


# ------------------------------------------------ optional checks that failed

def test_an_ai_check_that_raised_or_could_not_run_is_stored_as_could_not_be_checked(world, monkeypatch):
    run_id, _job = _enqueue(world)
    monkeypatch.setattr(settings, "review_ai_check_enabled", True)

    def boom(*a, **k):
        raise RuntimeError("document text that must not be stored")
    monkeypatch.setattr(ai_engineering_check, "run_check", boom)
    review_jobs._ai_check(run_id, world["scope"], [])
    status = ai_engineering_check.ai_check_status(run_id, allowed_document_ids=world["scope"])
    assert status["complete"] is False and status["ran"] is False
    assert "could not be checked" in status["plain"] and "RuntimeError" in status["plain"]
    assert "document text" not in json.dumps(status)

    # Asked for, and the lane is not available: also a fact on the run.
    monkeypatch.undo()
    monkeypatch.setattr(ai_engineering_check, "available", lambda: (False, "the Claude lane is off"))
    ai_engineering_check.run_check(run_id, allowed_document_ids=world["scope"], cited=[])
    again = ai_engineering_check.ai_check_status(run_id, allowed_document_ids=world["scope"])
    assert again["complete"] is False and "the Claude lane is off" in again["plain"]


def test_a_web_check_that_raised_is_stored_as_could_not_be_checked(world, monkeypatch):
    run_id, _job = _enqueue(world)
    monkeypatch.setattr(settings, "review_web_standards_enabled", True)

    def boom(*a, **k):
        raise RuntimeError("fetch failed with a private url")
    monkeypatch.setattr(web_standards, "run_check", boom)
    review_jobs._web_check(run_id, world["scope"], ["API 682"])
    status = main_mod._json_column(_run_row(run_id)["web_check_status"])
    assert status["complete"] is False
    assert "Web standards check could not be checked" in status["plain"]
    assert "private url" not in json.dumps(status)
    assert main_mod._json_column(None) is None            # never ran: nothing stored


# ----------------------------------------- the review code: large unchecked parts

CODES = (comparison.CODE_APPROVED, comparison.CODE_APPROVED_WITH_COMMENTS,
         comparison.CODE_REJECTED, comparison.CODE_MANUAL)
SUFFICIENT = {"sufficient": True, "fields_read": 100, "fields_estimated": 100,
              "extraction_coverage": 1.0}


def test_a_run_with_large_unchecked_parts_cannot_be_approved():
    met = [{"compliance_status": comparison.COMPLIANT}] * 3
    small = comparison.recommend_code(
        met, SUFFICIENT, unchecked_counts={"not_compared": 1, "not_applied": 1, "checked": 3})
    assert small["code"] == comparison.CODE_APPROVED                       # control
    big = comparison.recommend_code(
        met, SUFFICIENT, unchecked_counts={"not_compared": 900, "not_applied": 100, "checked": 3})
    assert big["code"] == comparison.CODE_MANUAL
    assert "900 of 903 requirements that apply to this submittal were not checked" in big["reason"]
    # #678: requirements that DO NOT APPLY never weigh against a run, however many.
    inapplicable = comparison.recommend_code(
        met, SUFFICIENT, unchecked_counts={"not_compared": 1, "not_applied": 5000, "checked": 3})
    assert inapplicable["code"] == comparison.CODE_APPROVED
    # Approved with comments is an approval too.
    with_comments = comparison.recommend_code(
        met + [{"compliance_status": comparison.MISSING_INFORMATION}], SUFFICIENT,
        unchecked_counts={"not_compared": 50, "not_applied": 0, "checked": 4})
    assert with_comments["code"] == comparison.CODE_MANUAL
    # A proven breach is still a rejection: unchecked parts do not soften it.
    breach = comparison.recommend_code(
        [{"compliance_status": comparison.NON_COMPLIANT}], SUFFICIENT,
        unchecked_counts={"not_compared": 900, "not_applied": 0, "checked": 1})
    assert breach["code"] == comparison.CODE_REJECTED
    assert absence.unchecked_share(0, 0) is None


def test_the_run_hands_its_unchecked_counts_to_the_review_code(world, monkeypatch):
    seen = {}
    real = comparison.recommend_code

    def spy(findings, completeness, **kwargs):
        seen.update(kwargs.get("unchecked_counts") or {"missing": True})
        return real(findings, completeness, **kwargs)

    monkeypatch.setattr(comparison, "recommend_code", spy)
    run_id, job_id = _enqueue(world)
    review_jobs.claim_next(W)
    assert review_jobs.run(job_id, W) == "done"
    assert set(seen) == {"not_compared", "not_applied", "checked"}
    assert seen["checked"] >= 1                         # the noise requirement was checked


def test_a_real_run_stores_what_it_left_unchecked_and_feeds_the_code(world):
    run_id, job_id = _enqueue(world)
    review_jobs.claim_next(W)
    assert review_jobs.run(job_id, W) == "done"
    stored = json.loads(_run_row(run_id)["refusal_reason"])
    assert "table_values_not_compared" in stored and "requirements_not_applied" in stored
    assert stored["standards_not_checked"] == []


# --------------------------------------------- datasheet self-checks: not silent

def _fact(name, value, unit=None, page=1, **kw):
    return {"id": f"f-{name}", "field_name": name, "field_label": name, "raw_value": value,
            "field_value": value, "raw_unit": unit, "page": page, "is_blank": 0,
            "equipment_tag": None, "normalized_value": None, "normalized_unit": None, **kw}


def test_a_revision_block_check_that_could_not_run_says_so_instead_of_vanishing():
    assert datasheet_checks.revision_check_not_run("Pressure Vessel", {}) == \
        "no page text was read for this datasheet"
    assert datasheet_checks.revision_check_not_run(None, {1: "text"}) == \
        "the equipment type is unknown"
    # Control: with page text and a type the check runs.
    assert datasheet_checks.revision_check_not_run("Pressure Vessel", {1: "text"}) is None
    ran = datasheet_checks.evaluate([], equipment_type="Pressure Vessel", page_texts={1: "text"})
    assert [r["status"] for r in ran if r["rule_id"] == "DS-R1"] == [datasheet_checks.MISSING_INFORMATION]
    # And the export lists it among the parts that could not be checked.
    parts = absence.unchecked_parts(
        run_status="completed",
        outcome={"datasheet_check_not_run": "the equipment type is unknown"})
    assert [p["part"] for p in parts] == ["datasheet_check_not_run", "unchecked_share"]
    assert "could not be checked: the equipment type is unknown" in parts[0]["line"]


def test_a_missing_revision_block_is_not_the_contractors_omission_while_a_page_is_unread(world):
    run_id, _job = _enqueue(world)
    result = datasheet_checks._result("DS-R1", datasheet_checks.MISSING_INFORMATION,
                                      "A datasheet carries a revision block.",
                                      "No revision block was found on the datasheet pages read.",
                                      None, "revision_block", None)
    unread = {"pages_total": 4, "fact_pages": [1], "pages_not_read_into_fields": [2, 3, 4],
              "pages_read_only_by_page_reader": []}
    [written] = datasheet_checks.store(run_id, world["sub"], [result], pages_read=unread)
    assert written["compliance_status"] == comparison.NEEDS_ENGINEER_REVIEW
    assert "2-4" in written["ai_rationale"]
    # An engineer-to-check finding says so; it never reads "None.".
    check = datasheet_checks._result("DS-R1", datasheet_checks.NEEDS_ENGINEER_REVIEW, "t", "d", None, "x", None)
    [written2] = datasheet_checks.store(run_id, world["sub"], [check], pages_read=unread)
    assert written2["required_action"] == "Engineer to check."


# ------------------------------------------------------------------ PDF export

def test_a_long_finding_is_not_cut_off_in_the_pdf_and_an_empty_report_is_not_a_pass(world):
    import pymupdf
    tail = "THE-END-MARKER-OF-A-VERY-LONG-FINDING"
    review.create({"document_id": world["sub"], "category": "technical_query", "severity": "minor",
                   "requirement": "r", "finding": ("word " * 700) + tail, "required_action": "a",
                   "status": "open", "approval_status": "pending"}, created_by=None)
    path = review.render_report(world["sub"], allowed_document_ids=None)
    text = " ".join(page.get_text() for page in pymupdf.open(str(path)))
    assert tail in text
    # No findings: the report says it is NOT a statement of acceptance.
    other = _doc("sub_other", "CONTRACTOR_SUBMITTAL", "nothing here")
    empty = " ".join(page.get_text() for page in pymupdf.open(
        str(review.render_report(other, allowed_document_ids=None))))
    assert "No review findings have been recorded" in empty
    assert "NOT a statement that the document was reviewed and found acceptable" in empty


# ---------------------------------------------------------------------- analysis

def test_a_named_document_that_cannot_be_read_is_listed_and_a_readable_one_is_not(world):
    scope = _scope_of(world)
    assert analysis.unresolved_names(scope, ["sub_pump.pdf", "doc16.pdf"]) == ["doc16.pdf"]
    assert analysis.unresolved_names(scope, []) == []


# ----------------------------------------------- the unchecked share, as a percentage

def test_every_incomplete_sheet_shows_the_unchecked_share_with_its_denominator():
    outcome = {"unchecked_counts": {"not_compared": 30, "not_applied": 10, "checked": 40}}
    parts = absence.unchecked_parts(run_status="completed", outcome=outcome)
    assert [p["part"] for p in parts] == ["unchecked_share", "requirement_reasons", "requirement_reasons"]
    assert ("Of 80 requirements in scope: 40 checked, 30 apply but were not checked, "
            "10 do not apply. 30 of 70 (43%) of the requirements that apply were not checked"
            ) in parts[0]["line"]
    assert "30 of 70 (43%)" in absence.notice_for(parts)
    # A run that never reached the comparison says the share is not known, never 0%.
    failed = absence.unchecked_parts(run_status="failed", outcome={})
    assert failed[-1]["part"] == "unchecked_share" and "not known" in failed[-1]["line"]
    assert "0%" not in failed[-1]["line"]
    # Nothing unchecked, nothing to show.
    assert absence.unchecked_parts(run_status="completed", outcome={
        "unchecked_counts": {"not_compared": 0, "not_applied": 0, "checked": 9}}) == []


def test_the_run_stores_the_counts_the_share_is_taken_from(world):
    run_id, _job = _enqueue(world)
    comparison._store_run_outcome(
        run_id, {"code": comparison.CODE_MANUAL, "reason": "r"}, {},
        unchecked_counts={"not_compared": 3, "not_applied": 1, "checked": 4})
    row = db.connect().execute("SELECT refusal_reason FROM review_runs WHERE id = ?", (run_id,)).fetchone()
    assert json.loads(row["refusal_reason"])["unchecked_counts"] == {
        "not_compared": 3, "not_applied": 1, "checked": 4}


def test_the_limit_is_read_from_the_data_file_and_a_bad_value_never_means_no_limit(
        tmp_path, monkeypatch):
    path = tmp_path / "thresholds.json"
    monkeypatch.setattr(absence, "THRESHOLDS_PATH", path)
    path.write_text('{"unchecked_share_limit": 0.8}', encoding="utf-8")
    assert absence.unchecked_share_limit() == 0.8
    # The shipped file says 0.5 (owner decision 2026-10-09).
    monkeypatch.undo()
    assert absence.unchecked_share_limit() == 0.5
    monkeypatch.setattr(absence, "THRESHOLDS_PATH", path)
    for bad in ('{"unchecked_share_limit": 0}', '{"unchecked_share_limit": "x"}', "not json", "{}"):
        path.write_text(bad, encoding="utf-8")
        assert absence.unchecked_share_limit() == 0.5
    # And the review code obeys the file: at 0.9 a 60% unchecked run is not forced to manual review.
    path.write_text('{"unchecked_share_limit": 0.9}', encoding="utf-8")
    met = [{"compliance_status": comparison.COMPLIANT}] * 4
    code = comparison.recommend_code(
        met, SUFFICIENT, unchecked_counts={"not_compared": 6, "not_applied": 0, "checked": 4})["code"]
    assert code != comparison.CODE_MANUAL
    path.write_text('{"unchecked_share_limit": 0.5}', encoding="utf-8")
    assert comparison.recommend_code(
        met, SUFFICIENT, unchecked_counts={"not_compared": 6, "not_applied": 0, "checked": 4}
    )["code"] == comparison.CODE_MANUAL
