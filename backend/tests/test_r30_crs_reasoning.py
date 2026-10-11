"""R30 (#763): every CRS row of the INTERNAL copy shows the chain behind it -
the value read and its page, the clause, the result, and the code's own why
(a calculation's inputs and numbers included). The issue copy never does.

What must hold: the chain is built only from what the findings stored (a
value not read says so; nothing stored prints nothing); every machine row
and AI/web item carries it; the internal view and workbook print it in a new
last column; the issue copy has neither the column nor the text.
Invented findings only.
"""
from __future__ import annotations

import io

import openpyxl

from app import crs_export, crs_mapping

BREACH = {"id": "f1", "compliance_status": "NON_COMPLIANT", "contractor_section": "Noise level",
          "contractor_evidence_text": "95 dB(A)", "contractor_page": 2,
          "standard_name": "API 610.pdf", "standard_clause": "6.1.14", "standard_page": 12,
          "ai_rationale": "95 dB(A) is above the limit of 90 dB(A).", "requirement": "Noise",
          "finding": "Noise too high", "required_action": "Reduce"}
CALC = {"id": "f2", "origin": "datasheet_check", "compliance_status": "NON_COMPLIANT",
        "contractor_section": "Effective orifice area", "contractor_evidence_text": "2.853 in2",
        "contractor_page": 3, "requirement_source_text": "Datasheet check PSV-C01: ...",
        "ai_rationale": "Datasheet check PSV-C01: Effective orifice area 2.853 in2 (page 3) is below "
                        "the required effective area 3,699.0 mm2 calculated by API 520-I 5.6.3",
        "finding": "Datasheet check: below", "required_action": "Contractor to correct the datasheet."}


def test_the_chain_is_value_page_clause_result_and_why():
    assert crs_mapping.reasoning_text([BREACH]) == (
        "Read: Noise level = 95 dB(A) (p.2). Clause: API 610 cl. 6.1.14 (p.12). "
        "Result: does not meet. Why: 95 dB(A) is above the limit of 90 dB(A).")


def test_a_value_not_read_says_so_and_nothing_stored_prints_nothing():
    blank = {**BREACH, "contractor_evidence_text": None, "compliance_status": "MISSING_INFORMATION"}
    assert crs_mapping.reasoning_text([blank]).startswith("Read: Noise level (no value read) (p.2).")
    assert "Result: value not given" in crs_mapping.reasoning_text([blank])
    assert crs_mapping.reasoning_text([{"id": "x"}]) == ""


def test_a_calculation_row_carries_its_numbers():
    rows = crs_mapping.build_crs_rows([CALC], [], "psv.pdf")
    row = next(r for r in rows if r["finding_id"] == "f2")
    assert "calculated by API 520-I 5.6.3" in row["reasoning"]
    assert "Clause: Datasheet check PSV-C01" in row["reasoning"]


def test_every_machine_row_carries_its_chain():
    rows = crs_mapping.build_crs_rows([BREACH, CALC], [], "sheet.pdf")
    machine = [r for r in rows if r.get("finding_id")]
    assert machine and all(r["reasoning"] for r in machine)


def _book(copy):
    rows = crs_mapping.build_crs_rows([BREACH], [], "sheet.pdf")
    for r in rows:
        r["engineer_confirmed"] = True      # so the issue copy keeps the row
    meta = {"review_run_id": "run-1", "copy": copy}
    view = crs_export.build_crs_view(rows, meta)
    ws = openpyxl.load_workbook(io.BytesIO(crs_export.build_crs(rows, meta))).active
    return view, ws


def test_the_internal_copy_prints_it_in_a_new_last_column():
    view, ws = _book("internal")
    assert view["columns"][-1] == crs_export.REASONING_COLUMN
    assert view["columns"][-2] == crs_export.AI_COLUMN          # nothing moved
    assert ws.cell(row=crs_export.COLUMN_HEADER_ROW, column=11).value == crs_export.REASONING_COLUMN
    assert ws.cell(row=crs_export.COLUMN_HEADER_ROW + 1, column=11).value == view["rows"][0]["reasoning"]
    assert view["rows"][0]["reasoning"].startswith("Read: Noise level = 95 dB(A)")


def test_the_issue_copy_never_carries_it():
    view, ws = _book("issue")
    assert crs_export.REASONING_COLUMN not in view["columns"]
    assert view["rows"] and all(r["reasoning"] == "" for r in view["rows"])
    header = [ws.cell(row=crs_export.COLUMN_HEADER_ROW, column=c).value for c in range(1, 13)]
    assert crs_export.REASONING_COLUMN not in header
    assert ws.cell(row=crs_export.COLUMN_HEADER_ROW + 1, column=11).value in (None, "")
