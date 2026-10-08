"""Permanent CRS comment numbers and the reviewer's Open/Closed status.

Industry practice for a Comment Resolution Sheet (researched 2026-09-29): a
comment's ID is permanent, never reused, and follows it to the next revision;
only the reviewer closes a comment. Here:

  * a comment gets "CRS-<submittal no>-001" in the client's own Item No column
    the moment an engineer makes it theirs - and the export and preview,
    which must write nothing, never mint one;
  * the number survives a re-export and a re-run;
  * a number is never handed out twice, and a confirmed comment always keeps
    its number over a same-subject machine draft written beside it;
  * Final Resolution reads "Open" until a signed-in reviewer closes it.

Synthetic sheet only (no client text), via `test_review_screen`'s world.
"""
from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import comparison, crs_export, crs_mapping, crs_numbers, db
from app.main import app
from tests.test_b3_page_ledger import temp_storage  # noqa: F401 - autouse
from tests.test_model_matching import _signed_in
from tests.test_review_screen import _first_comment, _noise_rows, world  # noqa: F401


def _count() -> int:
    return db.connect().execute("SELECT COUNT(*) FROM crs_comment_numbers").fetchone()[0]


def _confirm(run, scope, monkeypatch) -> str:
    finding_id, _ = _first_comment(run)
    _signed_in(monkeypatch, scope)
    response = TestClient(app).patch(f"/api/reviews/findings/{finding_id}",
                                     json={"confirmed": True})
    assert response.status_code == 200, response.text
    return finding_id


# ============================================================ the routes

def test_the_export_and_preview_never_mint_a_number(world):
    """They write nothing - the project's own rule for them - so an
    unconfirmed machine comment stays unnumbered however often it is viewed."""
    _sub, run, _scope = world
    client = TestClient(app)
    assert client.get(f"/api/reviews/runs/{run}/crs/preview").status_code == 200
    assert client.get(f"/api/reviews/runs/{run}/crs").status_code == 200
    assert _count() == 0
    [row] = _noise_rows(run)
    assert isinstance(row["item_no"], int) and row["crs_ref"] == ""
    assert row["final_resolution"] == ""


def test_confirming_a_comment_gives_it_a_permanent_number(world, monkeypatch):
    _sub, run, scope = world
    _confirm(run, scope, monkeypatch)
    [row] = _noise_rows(run)
    assert row["item_no"] == row["crs_ref"]
    assert row["crs_ref"].startswith("CRS-") and row["crs_ref"].endswith("-001")
    assert row["final_resolution"] == "Open"
    assert not row["comment"].startswith("Ref:"), "one ID per comment, not two"

    workbook = openpyxl.load_workbook(io.BytesIO(
        TestClient(app).get(f"/api/reviews/runs/{run}/crs").content))
    column_a = [c.value for c in workbook.active["A"]]
    assert row["crs_ref"] in column_a, "the xlsx prints the same number as the preview"


def test_the_number_survives_a_re_export_and_a_re_run(world, monkeypatch):
    _sub, run, scope = world
    _confirm(run, scope, monkeypatch)
    [first] = _noise_rows(run)
    [again] = _noise_rows(run)
    assert again["crs_ref"] == first["crs_ref"]

    # A re-run keeps the confirmed comment AND writes a fresh machine proposal
    # about the same field beside it. The confirmed one keeps the number.
    #
    # THE CONFIRMATION IS MADE OLDER THAN THE RE-RUN, as it is in real use.
    # Timestamps have one-second resolution, and this test used to confirm and
    # re-run inside the same second: the tie left the order to chance, so it
    # passed about 19 times in 20 while every real re-run - always a second or
    # more later - printed the draft's "AI Review" (honesty audit entry 81).
    db.connect().execute(
        "UPDATE review_findings SET updated_at = '2000-01-01T00:00:00Z'"
        " WHERE confirmed_by IS NOT NULL")
    db.connect().commit()
    comparison.run_comparison(run, allowed_document_ids=scope)
    rows = _noise_rows(run)
    numbered = [r for r in rows if r["crs_ref"]]
    assert [r["crs_ref"] for r in numbered] == [first["crs_ref"]]
    assert numbered[0]["comment_by"] == first["comment_by"]


def test_a_reviewer_closes_and_reopens_a_comment_under_their_own_name(world, monkeypatch):
    _sub, run, scope = world
    _confirm(run, scope, monkeypatch)
    ref = _noise_rows(run)[0]["crs_ref"]
    client = TestClient(app)

    closed = client.post(f"/api/reviews/runs/{run}/crs/comments/{ref}/status",
                         json={"status": "Closed"})
    assert closed.status_code == 200, closed.text
    assert closed.json()["status"] == "Closed" and closed.json()["status_by"] == "eng-1"
    assert _noise_rows(run)[0]["final_resolution"] == "Closed"

    reopened = client.post(f"/api/reviews/runs/{run}/crs/comments/{ref}/status",
                           json={"status": "Open"})
    assert reopened.status_code == 200
    assert _noise_rows(run)[0]["final_resolution"] == "Open"


def test_a_closure_nobody_signed_is_refused(world, monkeypatch):
    _sub, run, scope = world
    _confirm(run, scope, monkeypatch)
    ref = _noise_rows(run)[0]["crs_ref"]
    monkeypatch.setattr("app.config.settings.auth_mode", "disabled")
    response = TestClient(app).post(f"/api/reviews/runs/{run}/crs/comments/{ref}/status",
                                    json={"status": "Closed"})
    assert response.status_code == 401
    assert _noise_rows(run)[0]["final_resolution"] == "Open"


def test_a_number_that_is_not_this_submittals_is_404(world, monkeypatch):
    _sub, run, scope = world
    _confirm(run, scope, monkeypatch)
    ref = _noise_rows(run)[0]["crs_ref"]
    label = crs_numbers.parse_ref(ref)[0]
    client = TestClient(app)
    for bad in ("CRS-SOMEONEELSE-001", f"CRS-{label}-999", "not-a-number"):
        response = client.post(f"/api/reviews/runs/{run}/crs/comments/{bad}/status",
                               json={"status": "Closed"})
        assert response.status_code == 404, bad


# ============================================================ the rules

def test_numbers_are_sequential_idempotent_and_never_reused():
    scope, label = crs_numbers.scope_for("SUB-7", "doc_a")
    got = crs_numbers.assign(scope, label, ["k1", "k2"], document_id="doc_a", review_run_id="r1")
    assert [got["k1"]["ref"], got["k2"]["ref"]] == ["CRS-SUB-7-001", "CRS-SUB-7-002"]
    # Asking again changes nothing.
    assert crs_numbers.assign(scope, label, ["k2", "k1"], document_id="doc_a",
                              review_run_id="r2")["k1"]["seq"] == 1
    # "k2" gone from a later sheet does not free 002: the next comment is 003.
    got = crs_numbers.assign(scope, label, ["k1", "k3"], document_id="doc_a", review_run_id="r3")
    assert got["k3"]["seq"] == 3


def test_a_revision_with_the_same_submittal_number_continues_one_sequence():
    first = crs_numbers.scope_for("KJO-SUB-0042", "doc_rev0")
    second = crs_numbers.scope_for(" kjo-sub-0042 ", "doc_rev1")
    assert first == second
    crs_numbers.assign(*first, ["a"], document_id="doc_rev0", review_run_id="r")
    got = crs_numbers.assign(*second, ["b"], document_id="doc_rev1", review_run_id="r")
    assert got["b"]["ref"] == "CRS-KJO-SUB-0042-002"


def test_a_submittal_with_no_number_is_sequenced_by_its_document_not_a_guess():
    scope_a, label_a = crs_numbers.scope_for("", "doc_1a2b3c4d5e6f")
    scope_b, _ = crs_numbers.scope_for(None, "doc_ffffeeee0000")
    assert scope_a != scope_b and label_a == "1A2B3C4D"


def test_a_confirmed_row_claims_the_key_over_a_same_subject_draft():
    rows = [{"comment_key": "k", "engineer_confirmed": False},
            {"comment_key": "k", "engineer_confirmed": True}]
    assert crs_numbers.row_keys(rows) == ["k#2", "k"]


def test_the_comment_key_is_the_subject_not_its_spacing_or_run():
    assert crs_mapping.comment_key("Noise  Level", "95 dB(A)") == \
        crs_mapping.comment_key("noise level", " 95 dB(A) ")
    assert crs_mapping.comment_key("noise level", "95 dB(A)") != \
        crs_mapping.comment_key("noise level", "85 dB(A)")


def test_a_text_item_number_is_written_on_the_rows_own_line():
    """The workbook used Item No as the sheet row. A text number must not be."""
    rows = [{"document_name": "d", "page_section": "p1", "comment": "one",
             "comment_by": "e", "crs_ref": "CRS-X-007", "crs_status": "Open"},
            {"document_name": "d", "page_section": "p2", "comment": "two",
             "comment_by": "e"}]
    sheet = openpyxl.load_workbook(io.BytesIO(crs_export.build_crs(rows, {}))).active
    first = crs_export.COLUMN_HEADER_ROW + 1
    assert sheet.cell(first, 1).value == "CRS-X-007"
    assert sheet.cell(first, 7).value == "Open"
    assert sheet.cell(first + 1, 1).value == 2
    assert str(sheet.cell(first + 1, 4).value).startswith("Ref: RF-")


# ================================================= after issue: the reply loop

def _ref(run) -> str:
    return _noise_rows(run)[0]["crs_ref"]


def test_a_recorded_reply_prints_code_and_words_and_is_attributed(world, monkeypatch):
    _sub, run, scope = world
    _confirm(run, scope, monkeypatch)
    ref = _ref(run)
    response = TestClient(app).post(
        f"/api/reviews/runs/{run}/crs/comments/{ref}/response",
        json={"code": "Rejected", "text": "Vendor rating governs."})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["response_by"] == "eng-1" and body["response_source"] == crs_numbers.SOURCE_RECORDED
    [row] = _noise_rows(run)
    assert row["contractor_response"] == "Rejected: Vendor rating governs."
    assert row["final_resolution"] == "Open", "a reply is not a closure"


def test_the_returned_sheet_imports_by_number_and_never_guesses_a_code(world, monkeypatch):
    _sub, run, scope = world
    _confirm(run, scope, monkeypatch)
    ref = _ref(run)
    label = crs_numbers.parse_ref(ref)[0]
    client = TestClient(app)

    # The contractor's copy: the sheet we issued, with their reply typed in,
    # plus rows that must NOT be matched by position or guessed.
    book = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/reviews/runs/{run}/crs").content))
    sheet = book.active
    first = crs_export.COLUMN_HEADER_ROW + 1
    assert sheet.cell(first, 1).value == ref
    sheet.cell(first, 6).value = "Accepted with comments - will revise on Rev 1"
    extra = first + 40
    for offset, (item, reply) in enumerate((
            ("CRS-SOMEONEELSE-001", "Accepted"),
            (f"CRS-{label}-999", "Accepted"),
            ("7", "Accepted"))):
        sheet.cell(extra + offset, 1).value = item
        sheet.cell(extra + offset, 6).value = reply
    buffer = io.BytesIO()
    book.save(buffer)

    result = client.post(f"/api/reviews/runs/{run}/crs/reply",
                         files={"file": ("reply.xlsx", buffer.getvalue(),
                                         "application/octet-stream")})
    assert result.status_code == 200, result.text
    body = result.json()
    assert body["updated"] == 1
    assert body["other_submittal"] == 1 and body["unknown_number"] == 1
    assert body["not_a_crs_number"] >= 1
    # Every row with an Item No is in exactly one count.
    assert body["rows_read"] == sum(body[k] for k in (
        "updated", "updated_without_code", "no_response", "not_a_crs_number",
        "other_submittal", "unknown_number"))
    [row] = _noise_rows(run)
    assert row["contractor_response"] == "Accepted with comment: will revise on Rev 1"


def test_a_reply_that_states_no_code_is_stored_without_one():
    code, text = crs_numbers.parse_response("We will revise the datasheet.")
    assert code is None and text == "We will revise the datasheet."
    assert crs_numbers.parse_response("Rejected - see p.4") == ("Rejected", "see p.4")
    assert crs_numbers.parse_response("Accepted with comments") == ("Accepted with comment", "")


def test_a_file_that_is_not_a_crs_is_refused(world, monkeypatch):
    _sub, run, scope = world
    _confirm(run, scope, monkeypatch)
    response = TestClient(app).post(f"/api/reviews/runs/{run}/crs/reply",
                                    files={"file": ("x.xlsx", b"not a workbook", "text/plain")})
    assert response.status_code == 422


def test_an_open_comment_is_carried_forward_until_a_reviewer_closes_it(world, monkeypatch):
    """A run that no longer raises an issued comment must not drop it.

    CHANGED 2026-09-30 (audit): this test used to reject a comment that was
    never issued and assert it was still printed as carried forward - the
    defect itself (a rejected comment must never be issued). The comment is
    now ISSUED first - the contractor's reply is on record - so it stays on
    the sheet after the rejection until a reviewer closes it. The unissued
    case is `test_audit_crs_fixes.test_a_confirmed_then_rejected_comment_*`.
    """
    _sub, run, scope = world
    finding_id = _confirm(run, scope, monkeypatch)
    ref = _ref(run)
    client = TestClient(app)
    assert client.post(f"/api/reviews/runs/{run}/crs/comments/{ref}/response",
                       json={"code": "Rejected", "text": "we disagree"}).status_code == 200
    # The finding stops being raised (here: rejected, so the mapping drops it).
    assert client.patch(f"/api/reviews/findings/{finding_id}",
                        json={"approval_status": "rejected"}).status_code == 200
    rows = client.get(f"/api/reviews/runs/{run}/crs/preview").json()["rows"]
    carried = [r for r in rows if r["crs_ref"] == ref]
    assert len(carried) == 1 and carried[0]["row_kind"] == "carried_forward"
    assert carried[0]["final_resolution"] == "Open"
    assert "95 dB(A)" in carried[0]["comment"], "it prints what the comment last said"
    issue = client.get(f"/api/reviews/runs/{run}/crs/preview?copy=issue")
    assert issue.status_code == 200
    assert ref in [r["crs_ref"] for r in issue.json()["rows"]], "an issued comment stays issued"

    assert client.post(f"/api/reviews/runs/{run}/crs/comments/{ref}/status",
                       json={"status": "Closed", "note": "withdrawn"}).status_code == 200
    rows = client.get(f"/api/reviews/runs/{run}/crs/preview").json()["rows"]
    assert ref not in [r["crs_ref"] for r in rows], "a closed comment is not carried"


def test_the_closing_note_prints_and_the_history_says_who_did_what(world, monkeypatch):
    _sub, run, scope = world
    _confirm(run, scope, monkeypatch)
    ref = _ref(run)
    client = TestClient(app)
    client.post(f"/api/reviews/runs/{run}/crs/comments/{ref}/response",
                json={"code": "Accepted", "text": ""})
    client.post(f"/api/reviews/runs/{run}/crs/comments/{ref}/status",
                json={"status": "Closed", "note": "verified on Rev 1 p.4"})
    [row] = _noise_rows(run)
    assert row["final_resolution"] == "Closed: verified on Rev 1 p.4"

    history = client.get(f"/api/reviews/runs/{run}/crs/comments/{ref}/history")
    assert history.status_code == 200
    events = [(e["event"], e["by"]) for e in history.json()["events"]]
    assert [e for e, _by in events] == ["numbered", "response", "closed"]
    # Names, not ids: `_signed_in` names its engineer "Engineer".
    assert all(by == "Engineer" for _e, by in events)


def test_the_issued_sheet_asks_for_a_response_code_on_every_reply_cell():
    rows = [{"document_name": "d", "page_section": "p", "comment": "c",
             "comment_by": "e", "crs_ref": "CRS-X-001"}] * 3
    sheet = openpyxl.load_workbook(io.BytesIO(crs_export.build_crs(rows, {}))).active
    [tip] = sheet.data_validations.dataValidation
    first = crs_export.COLUMN_HEADER_ROW + 1
    assert str(tip.sqref) == f"F{first}:F{first + 2}"
    assert "Accepted" in tip.prompt and "Clarification needed" in tip.prompt
    assert tip.showErrorMessage is False, "a prompt, never a restriction"


def test_a_numbered_comment_raised_again_unconfirmed_is_still_issued(world, monkeypatch):
    """A resubmittal, or a re-run that writes a fresh finding, raises the SAME
    comment (same subject, same stated value) without the engineer's
    confirmation on the new finding. It already holds a number, which is only
    ever minted for a comment an engineer confirmed - so it is issued, under
    that number, rather than silently dropped from the contractor's copy."""
    _sub, run, scope = world
    finding_id = _confirm(run, scope, monkeypatch)
    ref = _ref(run)
    with db.connect() as conn:  # the new finding: same subject, not confirmed
        conn.execute("UPDATE review_findings SET confirmed_by = NULL, confirmed_at = NULL"
                     " WHERE id = ?", (finding_id,))
    issue = TestClient(app).get(f"/api/reviews/runs/{run}/crs/preview?copy=issue").json()
    assert ref in [r["crs_ref"] for r in issue["rows"]]


# ============================================ numbered by the DOCUMENT number

def _classify(doc_id: str, **fields) -> None:
    cols = ", ".join(fields)
    marks = ", ".join("?" for _ in fields)
    updates = ", ".join(f"{k} = excluded.{k}" for k in fields)
    with db.connect() as conn:
        conn.execute(
            f"INSERT INTO document_classification (document_id, suggested_by, {cols})"
            f" VALUES (?, 'classifier', {marks})"
            f" ON CONFLICT(document_id) DO UPDATE SET {updates}",
            (doc_id, *fields.values()))


def test_comments_are_numbered_by_the_document_number_never_the_transmittal(world, monkeypatch):
    """A transmittal number changes with every submission; the document number
    printed on the datasheet does not. Numbering by the transmittal would
    start a new sequence on every resubmittal and break carry-forward."""
    sub, run, scope = world
    _classify(sub, document_number="KJO-DS-0007", transmittal_number="T-0001")
    _confirm(run, scope, monkeypatch)
    assert _ref(run).startswith("CRS-KJO-DS-0007-")


def test_editing_the_document_number_later_never_renumbers_an_issued_comment(world, monkeypatch):
    sub, run, scope = world
    _classify(sub, document_number="KJO-DS-0007")
    _confirm(run, scope, monkeypatch)
    ref = _ref(run)
    _classify(sub, document_number="KJO-DS-9999")   # corrected after issue
    assert _ref(run) == ref, "the comment keeps the number the contractor was given"
    response = TestClient(app).post(f"/api/reviews/runs/{run}/crs/comments/{ref}/status",
                                    json={"status": "Closed"})
    assert response.status_code == 200, "and it can still be closed under it"


def test_a_confirmed_finding_leads_a_row_it_shares_with_a_newer_draft():
    """Order-independent: whichever of the two the caller lists first, the
    row speaks for the engineer's confirmed finding."""
    base = {"compliance_status": "NON_COMPLIANT", "requirement_source_text": "noise not above 90 dB(A)",
            "crs_field_label": "Noise", "contractor_evidence_text": "95 dB(A)", "contractor_page": 1,
            "fact_id": "f1", "requirement_id": "r1", "origin": None}
    draft = {**base, "id": "new-draft"}
    confirmed = {**base, "id": "old-confirmed", "confirmed_by": "eng-1", "confirmed_by_name": "Eng One"}
    for order in ([draft, confirmed], [confirmed, draft]):
        rows = [r for r in crs_mapping.build_crs_rows(order, [], "S.pdf")
                if r["row_kind"] == crs_mapping.ROW_KIND_NON_COMPLIANT]
        assert len(rows) == 1
        assert rows[0]["comment_by"] == "AI Review, confirmed by Eng One"
        assert rows[0]["finding_id"] == "old-confirmed"
