"""PF5 (#885): no unfinished CRS reaches a contractor.

Pass mark (owner, PROOF FIRST, 13 Oct): the contractor-copy export is refused
while any page is unread or any finding/question is undecided; the engineer
can resolve each open item with a recorded decision (who, when, reason)
without waiting for AI; the exact issued workbook is archived with its hash
and run id; the internal copy stays available.

Invented runs and findings only.
"""
from __future__ import annotations

import hashlib
import io
import json
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import access, crs_issue, db
from app.main import app
from tests.test_crs_endpoint import NOW, _client, _finding, _submittal, temp_db  # noqa: F401 - fixture


def _run(doc_id: str, *, unread: list[int] | None = None, total: int | None = 4,
         final_code: str | None = "Approved") -> str:
    run_id = str(uuid.uuid4())
    coverage = None if total is None else {
        "pages_total": total, "fact_pages": [p for p in range(1, total + 1) if p not in (unread or [])],
        "pages_not_read_into_fields": unread or [],
        "not_read_reasons": {str(p): "unreadable" for p in unread or []}}
    outcome = {"recommended_code": "Manual Review Required", "reason": "r",
               **({"page_coverage": coverage} if coverage else {})}
    with db.connect() as conn:
        conn.execute(
            "INSERT INTO review_runs (id,submittal_document_id,status,outcome_json,"
            "engineer_final_code,created_at,updated_at) VALUES (?,?,'completed',?,?,?,?)",
            (run_id, doc_id, json.dumps(outcome), final_code, NOW, NOW))
    return run_id


def _fid(run_id: str) -> str:
    return db.connect().execute("SELECT id FROM review_findings WHERE review_run_id = ?",
                                (run_id,)).fetchone()["id"]


def _gate(client, run_id):
    r = client.get(f"/api/reviews/runs/{run_id}/crs/issue-gate")
    assert r.status_code == 200, r.text
    return r.json()


def _issue(client, run_id):
    return client.post(f"/api/reviews/runs/{run_id}/crs/issue")


def test_an_undecided_finding_refuses_the_contractor_copy_and_names_it():
    doc = _submittal()
    run = _run(doc)
    _finding(doc, run, "NON_COMPLIANT")
    client = _client(doc)
    gate = _gate(client, run)
    assert gate["issuable"] is False and gate["unread_pages"] == []
    assert [f["finding_id"] for f in gate["open_findings"]] == [_fid(run)]
    r = _issue(client, run)
    assert r.status_code == 409 and r.json()["detail"]["code"] == "crs_not_finished"
    assert "1 finding(s) on the sheet have no engineer's decision" in r.json()["detail"]["message"]
    # the internal copy stays available
    assert client.get(f"/api/reviews/runs/{run}/crs").status_code == 200


def test_an_unread_page_refuses_it_and_a_run_with_no_page_record_too():
    doc = _submittal()
    client = _client(doc)
    run = _run(doc, unread=[3])
    gate = _gate(client, run)
    assert gate["unread_pages"] == [{"page": 3, "reason": "unreadable"}] and not gate["issuable"]
    assert _issue(client, run).status_code == 409
    # no page coverage on record is not "every page read"
    blind = _run(doc, total=None)
    assert crs_issue.NO_PAGE_COVERAGE in _gate(client, blind)["blockers"]
    assert _issue(client, blind).status_code == 409


def test_the_engineer_resolves_each_item_with_a_recorded_decision_then_it_is_issued_and_kept():
    doc = _submittal()
    run = _run(doc, unread=[3])
    _finding(doc, run, "NON_COMPLIANT")
    client = _client(doc)
    fid = _fid(run)

    page = client.post(f"/api/reviews/runs/{run}/crs/open-items/pages/3/decision",
                       json={"reason": "read page 3 by eye: a general note, no values"})
    assert page.status_code == 200, page.text
    assert (page.json()["decided_by"], page.json()["kind"], page.json()["subject"]) == ("eng", "page", "3")
    assert page.json()["decided_at"]
    accept = client.post(f"/api/reviews/runs/{run}/crs/open-items/findings/{fid}/decision",
                         json={"decision": "accept", "reason": "breach confirmed against clause 6.2.2"})
    assert accept.status_code == 200, accept.text
    row = dict(db.connect().execute("SELECT * FROM review_findings WHERE id = ?", (fid,)).fetchone())
    assert (row["approval_status"], row["confirmed_by"], row["approved_by"]) == ("accepted", "eng", "eng")

    assert _gate(client, run)["issuable"] is True
    issued = _issue(client, run)
    assert issued.status_code == 200, issued.text
    assert "issue-to-contractor" in issued.headers["content-disposition"]
    sha = hashlib.sha256(issued.content).hexdigest()
    assert issued.headers["x-issued-sha256"] == sha
    # THE EXACT ISSUED COPY IS KEPT, with its hash and run, and reads back byte for byte
    copies = client.get(f"/api/reviews/runs/{run}/crs/issued").json()["copies"]
    assert [(c["sha256"], c["review_run_id"], c["issued_by"]) for c in copies] == [(sha, run, "eng")]
    again = client.get(f"/api/reviews/runs/{run}/crs/issued/{copies[0]['id']}")
    assert again.status_code == 200 and again.content == issued.content
    # the accepted comment is on the contractor's copy
    text = " ".join(str(c.value) for row in load_workbook(io.BytesIO(issued.content)).active.iter_rows()
                    for c in row if c.value)
    assert "6.2.2" in text


def test_a_rejected_finding_is_decided_too_and_stays_off_the_sheet():
    doc = _submittal()
    run = _run(doc)
    _finding(doc, run, "NON_COMPLIANT")
    client = _client(doc)
    r = client.post(f"/api/reviews/runs/{run}/crs/open-items/findings/{_fid(run)}/decision",
                    json={"decision": "reject", "reason": "the value is gauge, the limit absolute"})
    assert r.status_code == 200
    assert _gate(client, run)["open_findings"] == []
    assert _issue(client, run).status_code == 200


def test_the_final_code_is_still_required():
    doc = _submittal()
    run = _run(doc, final_code=None)
    client = _client(doc)
    gate = _gate(client, run)
    assert gate["final_code_recorded"] is False and gate["issuable"] is False
    assert _issue(client, run).json()["detail"]["code"] == "code_not_decided"


def test_the_get_route_no_longer_hands_out_a_contractor_copy():
    doc = _submittal()
    run = _run(doc)
    r = _client(doc).get(f"/api/reviews/runs/{run}/crs", params={"copy": "issue"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "issue_through_post"


def test_decisions_need_a_reason_a_name_and_a_real_open_item():
    doc = _submittal()
    run = _run(doc, unread=[3])
    _finding(doc, run, "NON_COMPLIANT")
    client = _client(doc)
    assert client.post(f"/api/reviews/runs/{run}/crs/open-items/pages/3/decision",
                       json={"reason": ""}).status_code == 422
    # a page the run did read is not an open item
    assert client.post(f"/api/reviews/runs/{run}/crs/open-items/pages/2/decision",
                       json={"reason": "looked at it"}).status_code == 422
    assert client.post(f"/api/reviews/runs/{run}/crs/open-items/findings/{_fid(run)}/decision",
                       json={"reason": "no decision given"}).status_code == 422
    assert client.post(f"/api/reviews/runs/{run}/crs/open-items/findings/nope/decision",
                       json={"decision": "accept", "reason": "x y z"}).status_code == 404
    # an anonymous caller cannot decide
    app.dependency_overrides[access.current_scope] = lambda: access.AccessScope(
        user_id=None, allowed_document_ids=frozenset({doc}))
    assert TestClient(app).post(f"/api/reviews/runs/{run}/crs/open-items/pages/3/decision",
                                json={"reason": "read it"}).status_code == 401


def test_a_stranger_sees_no_gate_and_issues_nothing():
    doc = _submittal()
    run = _run(doc)
    stranger = _client("other-doc")
    assert stranger.get(f"/api/reviews/runs/{run}/crs/issue-gate").status_code == 404
    assert _issue(stranger, run).status_code == 404
    assert stranger.get(f"/api/reviews/runs/{run}/crs/issued").status_code == 404


def test_an_archived_copy_whose_bytes_changed_is_refused():
    doc = _submittal()
    run = _run(doc)
    client = _client(doc)
    assert _issue(client, run).status_code == 200
    copy = client.get(f"/api/reviews/runs/{run}/crs/issued").json()["copies"][0]
    path = db.connect().execute("SELECT path FROM crs_issued_copies WHERE id = ?",
                                (copy["id"],)).fetchone()["path"]
    with open(path, "ab") as f:
        f.write(b"tampered")
    r = client.get(f"/api/reviews/runs/{run}/crs/issued/{copy['id']}")
    assert r.status_code == 409 and "sha256" in r.json()["detail"]["message"]


def test_the_gate_and_the_lists_write_nothing():
    doc = _submittal()
    run = _run(doc)
    client = _client(doc)
    client.get(f"/api/reviews/runs/{run}/crs/issue-gate")
    client.get(f"/api/reviews/runs/{run}/crs/issued")
    tables = {r["name"] for r in db.connect().execute("SELECT name FROM sqlite_master")}
    assert "crs_issue_decisions" not in tables and "crs_issued_copies" not in tables
