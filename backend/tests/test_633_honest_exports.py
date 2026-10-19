"""#633 items 6 and 7: the review PDF and the CRS are honest about the run.

Item 6, the PDF: text is never dropped (a long finding continues on the next
page); the header prints the run code and status, the unread pages and the
missing standards (the readiness strip's own data); a run that did not
complete prints "Run <status>: <reason>" and never "No review findings have
been recorded" as if it had completed.

Item 7, the CRS: a run that did not complete is said on row 1 of the sheet,
with the same two counts; a completed run's sheet is unchanged.

Synthetic data only; no document text in logs.
"""
from __future__ import annotations

import io
import json
import secrets

import openpyxl
import pymupdf
import pytest

from app import (
    access,
    auth,
    crs_export,
    db,
    export_status,
    fts_index,
    pdf_flow,
    review,
    review_jobs,
    standards,
    submittal_review,
)
from app import main as main_mod
from app.config import settings

NOW = "2026-10-18T00:00:00Z"
W = "w-test"


@pytest.fixture(autouse=True)
def temp_storage(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "db_path", tmp_path / "t633.sqlite")
    monkeypatch.setattr(settings, "auth_mode", access.AUTH_REQUIRED)
    monkeypatch.setattr(settings, "auth_secret", secrets.token_urlsafe(48))
    monkeypatch.setattr(settings, "geometry_reader_enabled", False)
    monkeypatch.setattr(settings, "job_max_running", 1)
    db.reset_connection()
    db.init_db()
    submittal_review.ensure_schema()
    submittal_review.migrate_facts_to_per_document()
    fts_index.ensure_schema()
    access.set_user_resolver(auth.resolve_user_id)
    yield
    access.set_user_resolver(None)
    monkeypatch.setattr(settings, "auth_mode", access.AUTH_DISABLED)
    db.reset_connection()


def _doc(doc_id, role, text):
    with db.connect() as conn:
        conn.execute("INSERT INTO documents (id,filename,sha256,size_bytes,stored_path,status,"
                     "page_count,uploaded_at) VALUES (?,?,?,1,'/x','ready',1,?)",
                     (doc_id, f"{doc_id}.pdf", f"sha-{doc_id}", NOW))
        conn.execute("INSERT INTO document_classification (document_id,suggested_by,document_role)"
                     " VALUES (?,?,?)", (doc_id, "test", role))
        conn.execute("INSERT INTO chunks (id,document_id,filename,ordinal,page_start,page_end,section,"
                     "kind,text,token_count,content_hash,retrievable) VALUES (?,?,?,0,1,1,'1',"
                     "'prose',?,1,?,1)", (f"{doc_id}-c1", doc_id, f"{doc_id}.pdf", text, f"h-{doc_id}"))
    fts_index.index_document(doc_id)
    return doc_id


def _user(user_id, docs):
    with db.connect() as conn:
        conn.execute("INSERT INTO users (id,email,display_name,password_hash,is_active,created_at)"
                     " VALUES (?,?,?,'x',1,?)", (user_id, f"{user_id}@x", user_id, NOW))
        rid = f"r-{user_id}"
        conn.execute("INSERT INTO roles (id,name,description,kind,created_at) VALUES (?,?,'',?,?)",
                     (rid, rid, "discipline", NOW))
        conn.execute("INSERT INTO user_roles (user_id,role_id,granted_at) VALUES (?,?,?)", (user_id, rid, NOW))
        conn.execute("INSERT OR IGNORE INTO roles (id,name,description,kind,created_at)"
                     " VALUES ('role_admin','admin','','capability',?)", (NOW,))
        adm = conn.execute("SELECT id FROM roles WHERE name = 'admin'").fetchone()["id"]
        conn.execute("INSERT INTO user_roles (user_id,role_id,granted_at) VALUES (?,?,?)", (user_id, adm, NOW))
        for d in docs:
            conn.execute("INSERT INTO document_role_access (document_id,role_id,granted_at)"
                         " VALUES (?,?,?)", (d, rid, NOW))


@pytest.fixture
def world():
    """One held standard (API 610) and one cited but NOT held (API 682)."""
    std = _doc("std_api", "COMPANY_STANDARD", "Equipment noise level shall not exceed 90 dB(A).")
    sub = _doc("sub_pump", "CONTRACTOR_SUBMITTAL",
               "Pump shall comply with API 610 and API 682. Noise level 95 dB(A)")
    with db.connect() as conn:
        conn.execute("UPDATE document_classification SET document_number = 'API 610' WHERE document_id = ?", (std,))
    standards.extract_requirements(std, allowed_document_ids=frozenset({std}))
    _user("adm", [std, sub])
    scope = frozenset({std, sub})
    return {"std": std, "sub": sub, "scope": scope,
            "access": access.AccessScope(user_id="adm", allowed_document_ids=scope)}


def _failed_run(w, status="failed", reason="the model host went away", created_at=None):
    """A run that ended without completing, as `review_jobs.run` records one.
    `created_at` orders it against another run made in the same second."""
    run_id, _job = review_jobs.enqueue(w["sub"], allowed_document_ids=w["scope"], requested_by="adm")
    with db.connect() as conn:
        conn.execute("UPDATE review_runs SET status = ?, outcome_json = ?,"
                     " created_at = COALESCE(?, created_at) WHERE id = ?",
                     (status, json.dumps({"error": reason}), created_at, run_id))
    return run_id


def _completed_run(w):
    run_id, job_id = review_jobs.enqueue(w["sub"], allowed_document_ids=w["scope"], requested_by="adm")
    review_jobs.claim_next(W)
    assert review_jobs.run(job_id, W) == "done"
    assert submittal_review.get_review_run(run_id, allowed_document_ids=w["scope"])["status"] == "completed"
    return run_id


def _pdf_text(w) -> str:
    path = review.render_report(w["sub"], allowed_document_ids=w["scope"])
    with pymupdf.open(str(path)) as pdf:
        return "\n".join(page.get_text() for page in pdf)


def _sheet_cells(data: bytes, sheet="CRS") -> dict:
    ws = openpyxl.load_workbook(io.BytesIO(data))[sheet]
    return {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}


# ------------------------------------------------------------------ item 6: PDF

def test_a_failed_runs_pdf_prints_the_status_line_and_never_no_findings(world):
    run_id = _failed_run(world)
    text = _pdf_text(world)
    assert "Run failed: the model host went away" in text
    assert "No review findings have been recorded" not in text
    assert f"Review run: {run_id}" in text and "Status: failed" in text
    # A cancelled run is said in its own word; its reason is the stored one.
    cancelled = _failed_run(world, "cancelled", "cancelled by an engineer before it finished",
                            created_at="2999-01-01T00:00:00Z")
    text = _pdf_text(world)
    assert f"Review run: {cancelled}" in text          # the newest run is the one reported
    assert "Run cancelled: cancelled by an engineer before it finished" in text
    assert "No review findings have been recorded" not in text


def test_a_document_with_no_run_keeps_the_not_a_pass_wording(world):
    text = _pdf_text(world)
    assert "Review run: none recorded for this document" in text
    assert "No review findings have been recorded" in text
    assert "NOT a statement that the document was reviewed and found acceptable" in text
    assert "Run " not in text.replace("Review run", "")


def test_a_long_comment_is_fully_present_in_the_pdf(world):
    words = [f"w{i}" for i in range(3000)]              # ~ 150 lines: more than one page
    long_token = "X" * 400                              # wider than the page on its own
    review.create({"document_id": world["sub"], "category": "technical_query", "severity": "minor",
                   "requirement": "r", "finding": " ".join(words) + " " + long_token + " THE-END",
                   "required_action": "a", "status": "open", "approval_status": "pending"},
                  created_by=None)
    path = review.render_report(world["sub"], allowed_document_ids=world["scope"])
    with pymupdf.open(str(path)) as pdf:
        pages = [page.get_text() for page in pdf]
    text = " ".join(pages)
    flat = text.replace("\n", "")
    assert len(pages) >= 2
    assert "THE-END" in text and long_token in flat
    missing = [w for w in words if w not in text.split()]
    assert missing == []
    assert "NOT SHOWN HERE" not in text
    assert pdf_flow.CONTINUED in text


def test_wrap_never_makes_a_line_wider_than_the_width():
    width = 120.0
    lines = pdf_flow.wrap("short words then " + "Y" * 300 + " and a tail\n\nnext paragraph", width, 8.5)
    assert all(pymupdf.get_text_length(line, fontname="helv", fontsize=8.5) <= width for line in lines)
    assert "".join(lines).replace(" ", "") == ("short words then " + "Y" * 300
                                               + " and a tail next paragraph").replace(" ", "")
    assert "" in lines                                  # the paragraph break is kept


def test_the_pdf_prints_the_run_code_and_the_unread_and_missing_counts(world):
    run_id = _completed_run(world)
    text = _pdf_text(world)
    readiness = main_mod._readiness_payload(world["sub"], world["access"])
    assert readiness["standards_missing"] == ["API 682"]  # positive control for the data
    assert f"Review run: {run_id}" in text and "Status: completed" in text
    assert "Missing standards: 1 cited standard(s) not in the library (API 682)" in text
    unread = readiness["unread_pages"]
    assert f"Unread pages: {len(unread)} page(s) not read into fields" in text
    assert "Run completed" not in text                  # a completed run has no status line


# ------------------------------------------------------------------ item 7: CRS

def test_the_crs_of_an_incomplete_run_carries_the_marker_on_row_one(world):
    run_id = _failed_run(world)
    rows, meta, _name, _stamp = main_mod._crs_content(run_id, world["access"], "internal")
    note = meta["run_status_note"]
    assert note.startswith("REVIEW INCOMPLETE: run failed - the model host went away.")
    assert "Unread pages: 0 page(s) not read into fields." in note
    assert "Missing standards: 1 cited standard(s) not in the library (API 682)." in note
    for copy in ("internal", "issue"):
        cells = _sheet_cells(crs_export.build_crs(rows, {**meta, "copy": copy}))
        assert str(cells[(1, 1)]).startswith("REVIEW INCOMPLETE: run failed"), copy
        assert "API 682" in cells[(1, 1)]
        # The column headers did not move: the template's geometry is kept.
        assert cells[(crs_export.COLUMN_HEADER_ROW, 1)] == crs_export.HEADERS[0]
    # The preview says the same thing as the file.
    view = crs_export.build_crs_view(rows, {**meta, "copy": "issue"})
    assert view["run_status_note"] == note
    # The unread pages the note counts are the run's own stored pages.
    assert export_status.crs_header_note("failed", "x", [3, 5], []).startswith(
        "REVIEW INCOMPLETE: run failed - x. Unread pages: 2 page(s) not read into fields (3, 5).")
    assert export_status.crs_header_note("running", None, [], []) == (
        "REVIEW INCOMPLETE: run running - no reason was recorded. Unread pages: 0 page(s) "
        "not read into fields. Missing standards: 0 cited standard(s) not in the library.")


def test_the_crs_of_a_completed_run_is_unchanged(world):
    run_id = _completed_run(world)
    rows, meta, _name, _stamp = main_mod._crs_content(run_id, world["access"], "internal")
    assert meta["run_status_note"] == ""
    assert export_status.crs_header_note("completed", "", [1], ["API 682"]) == ""
    before = {k: v for k, v in meta.items() if k != "run_status_note"}   # the pre-#633 meta
    for copy in ("internal", "issue"):
        new = openpyxl.load_workbook(io.BytesIO(crs_export.build_crs(rows, {**meta, "copy": copy})))
        old = openpyxl.load_workbook(io.BytesIO(crs_export.build_crs(rows, {**before, "copy": copy})))
        assert new.sheetnames == old.sheetnames
        for name in new.sheetnames:
            a, b = new[name], old[name]
            assert [[c.value for c in row] for row in a.iter_rows()] == \
                   [[c.value for c in row] for row in b.iter_rows()], (copy, name)
            assert sorted(map(str, a.merged_cells.ranges)) == sorted(map(str, b.merged_cells.ranges))
            assert {r: d.height for r, d in a.row_dimensions.items()} == \
                   {r: d.height for r, d in b.row_dimensions.items()}
        cells = _sheet_cells(crs_export.build_crs(rows, {**meta, "copy": copy}))
        assert "INCOMPLETE: run" not in str(cells.get((1, 1)))
