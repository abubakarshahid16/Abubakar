"""#986: the engineer can decide EVERY open item of the issue gate.

Found on the live PSV (PF5 check, 17 Oct): rows with no finding behind them
(standard not cited / not applicable) answered 404 to the decide route, so the
contractor copy could never be issued; and a decision on a grouped row closed
only its first finding, so the row came back with the next one behind it.
Invented runs and findings only.
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

from app import db
from app import main as main_mod
from tests.test_885_crs_issue_gate import _gate, _issue, _run
from tests.test_crs_endpoint import _client, _finding, _submittal, temp_db  # noqa: F401 - fixture

CHECK = {"required_not_cited": [{"identifier": "API 520", "reason": "the sheet is a relief valve"}],
         "cited_not_applicable": [{"identifier": "ISO 9999", "reason": "it covers pumps"}]}


def _decide(client, run, subject, decision):
    return client.post(f"/api/reviews/runs/{run}/crs/open-items/findings/{subject}/decision",
                       json={"decision": decision, "reason": "checked against the register"})


def _text(content: bytes) -> str:
    return " ".join(str(c.value) for row in load_workbook(io.BytesIO(content)).active.iter_rows()
                    for c in row if c.value)


def test_a_row_with_no_finding_is_decided_by_its_key_then_the_copy_is_issued(monkeypatch):
    monkeypatch.setattr(main_mod.applicability_mod, "standards_check", lambda *a, **k: CHECK)
    doc = _submittal()
    run = _run(doc)
    client = _client(doc)
    gate = _gate(client, run)
    kinds = {f["row_kind"]: f["finding_id"] for f in gate["open_findings"]}
    assert set(kinds) == {"standard_not_cited", "standard_not_applicable"}
    assert _issue(client, run).status_code == 409

    accept = _decide(client, run, kinds["standard_not_cited"], "accept")
    assert accept.status_code == 200, accept.text
    assert (accept.json()["decided_by"], accept.json()["decision"]) == ("eng", "accept")
    assert _decide(client, run, kinds["standard_not_applicable"], "reject").status_code == 200

    assert _gate(client, run)["open_findings"] == []
    issued = _issue(client, run)
    assert issued.status_code == 200, issued.text
    text = _text(issued.content)
    assert "API 520" in text            # accepted: on the contractor's copy
    assert "ISO 9999" not in text       # rejected: withdrawn


def test_an_unknown_key_is_still_404(monkeypatch):
    monkeypatch.setattr(main_mod.applicability_mod, "standards_check", lambda *a, **k: CHECK)
    doc = _submittal()
    run = _run(doc)
    assert _decide(_client(doc), run, "not-a-row-key", "accept").status_code == 404


def test_one_decision_on_a_grouped_row_decides_every_finding_behind_it():
    doc = _submittal()
    run = _run(doc)
    _finding(doc, run, "NON_COMPLIANT", equipment_tag="TAG-A")
    _finding(doc, run, "NON_COMPLIANT", equipment_tag="TAG-B")
    client = _client(doc)
    gate = _gate(client, run)
    assert len(gate["open_findings"]) == 1          # one row for the two tags
    r = _decide(client, run, gate["open_findings"][0]["finding_id"], "accept")
    assert r.status_code == 200, r.text
    statuses = [row["approval_status"] for row in db.connect().execute(
        "SELECT approval_status FROM review_findings WHERE review_run_id = ?", (run,))]
    assert statuses == ["accepted", "accepted"]
    assert _gate(client, run)["open_findings"] == []
    assert _issue(client, run).status_code == 200
