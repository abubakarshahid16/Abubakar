"""#725 F7 (#735), part 1: a clean CRS.

  * ONE ROW PER PROBLEM: the same requirement, field and outcome for several
    tags is one comment listing every tag, each with its own value and page
    when they differ (it was one row per tag whenever a value or page
    differed: "the same comment repeated per tag");
  * A SYSTEM MISS IS NOT THE CONTRACTOR'S: a requirement no field read from
    the sheet was paired with is an engineer's question, never "Contractor to
    state ...";
  * a SEVERITY column (the worst in the group), in both copies;
  * rows within a group ordered by datasheet page, then clause;
  * the row's leader (whose key is the permanent number) is the same
    whatever order the findings arrive in.
Invented findings only.

Mutations: M7101-M7108 (scripts/mutations/f7_735_clean_crs.py).
"""
from __future__ import annotations

import io

import openpyxl

from app import crs_export
from app.crs_mapping import build_crs_rows

NC = {"compliance_status": "NON_COMPLIANT", "standard_name": "SPEC-X.pdf", "standard_clause": "5.1",
      "standard_page": 7, "requirement_source_text": "The set pressure shall not exceed 10 barg.",
      "crs_field_label": "SET PRESSURE", "severity": "major", "fact_id": "f"}


def _nc(tag, value, page, **over):
    return dict(NC, id=f"nc-{tag}", equipment_tag=tag, contractor_evidence_text=value,
                contractor_page=page, fact_id=f"f-{tag}", **over)


def test_one_problem_for_three_tags_is_one_row_with_each_tags_value():
    """THE MUTATION TARGET: three rows before."""
    rows = build_crs_rows([_nc("PSV-1", "12 barg", 2), _nc("PSV-2", "11 barg", 3),
                           _nc("PSV-3", "12 barg", 2)], [], "psv.pdf")
    assert len(rows) == 1
    row = rows[0]
    assert row["page_section"].startswith("p.2")
    assert all(t in row["page_section"] for t in ("PSV-1", "PSV-2", "PSV-3"))
    for tag, value in (("PSV-1", "12 barg"), ("PSV-2", "11 barg"), ("PSV-3", "12 barg")):
        assert f"{tag} states {value}" in row["comment"], tag
    assert "(p.3)" in row["comment"]


def test_a_requirement_no_field_was_paired_with_is_an_engineers_question():
    unpaired = {"compliance_status": "MISSING_INFORMATION", "id": "u", "standard_name": "SPEC-X.pdf",
                "standard_clause": "6.2", "requirement_source_text": "The lift shall be stated.",
                "crs_field_label": "LIFT", "ai_rationale": "no field read from the submittal answers this requirement"}
    [row] = build_crs_rows([unpaired], [], "psv.pdf")
    assert "Engineer to check whether the datasheet states LIFT" in row["comment"]
    assert "Contractor to" not in row["comment"]


def test_a_blank_the_contractor_left_is_still_the_contractors():
    blank = dict(NC, compliance_status="MISSING_INFORMATION", id="b", equipment_tag="PSV-1",
                 contractor_evidence_text="TBA", contractor_page=2, crs_is_blank=True, severity="minor")
    [row] = build_crs_rows([blank], [], "psv.pdf")
    assert "Contractor to provide" in row["comment"]


def test_a_row_carries_the_worst_severity_of_its_findings():
    rows = build_crs_rows([_nc("PSV-1", "12 barg", 2, severity="minor"),
                           _nc("PSV-2", "13 barg", 2, severity="major")], [], "psv.pdf")
    assert rows[0]["severity"] == "major"
    assert build_crs_rows([dict(_nc("PSV-1", "12 barg", 2), severity=None)], [], "p.pdf")[0]["severity"] == ""


def test_rows_in_a_group_are_ordered_by_page_then_clause():
    a = dict(NC, id="a", standard_clause="8.2", requirement_source_text="Rule A shall hold.",
             crs_field_label="A", contractor_page=2, equipment_tag="X", contractor_evidence_text="1")
    b = dict(NC, id="b", standard_clause="9.1", requirement_source_text="Rule B shall hold.",
             crs_field_label="B", contractor_page=1, equipment_tag="X", contractor_evidence_text="1")
    c = dict(NC, id="c", standard_clause="8.10", requirement_source_text="Rule C shall hold.",
             crs_field_label="C", contractor_page=2, equipment_tag="X", contractor_evidence_text="1")
    rows = build_crs_rows([a, b, c], [], "s.pdf")
    # page 1 first; then on page 2 clause 8.2 before 8.10 (numbers, not text)
    assert [r["finding_id"] for r in rows] == ["b", "a", "c"]


def test_the_rows_number_key_does_not_depend_on_input_order():
    findings = [_nc("PSV-2", "11 barg", 3), _nc("PSV-1", "12 barg", 2)]
    first = build_crs_rows(findings, [], "psv.pdf")[0]["comment_key"]
    again = build_crs_rows(list(reversed(findings)), [], "psv.pdf")[0]["comment_key"]
    assert first == again


def test_the_workbook_prints_the_severity_column_in_both_copies():
    rows = build_crs_rows([_nc("PSV-1", "12 barg", 2)], [], "psv.pdf")
    for copy in (crs_export.COPY_INTERNAL, crs_export.COPY_ISSUE):
        for r in rows:
            r["engineer_confirmed"] = True
        view = crs_export.build_crs_view(rows, {"copy": copy})
        assert "Severity" in view["columns"]
        book = openpyxl.load_workbook(io.BytesIO(crs_export.build_crs(rows, {"copy": copy})))
        ws = book["CRS"]
        assert ws.cell(row=crs_export.COLUMN_HEADER_ROW, column=9).value == "Severity"
        assert ws.cell(row=crs_export.COLUMN_HEADER_ROW + 1, column=9).value == "major"
