"""PF6 (#886): the engineer review pack - form only, proved on invented data.

Builds a sample of MAIN (non-PSV) findings and the full SECONDARY (PSV) set
into one sheet with a blank Mark/Note column per row, never a default; scores
a filled-in copy back into counts only (ids and numbers, never a finding's
text). Invented findings only; no PF1 run, no live data.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from tools import review_pack as rp

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "scripts"))
import review_pack as cli  # noqa: E402


def _finding(id_, standard="STD-1", clause="4.2", page=1, status="NON_COMPLIANT", text="UNIQUE TEXT"):
    return {"id": id_, "standard": standard, "clause": clause, "field": "design pressure",
            "compliance_status": status, "requirement_text": text, "submittal_page": page}


MAIN_FINDINGS = [_finding(f"m{n}") for n in range(1, 6)]
SECONDARY_FINDINGS = [_finding(f"s{n}", standard="PSV-STD") for n in range(1, 4)]


# ----------------------------------------------------------------- building

def test_sampling_is_reproducible_for_the_same_seed():
    first = rp.sample_rows(MAIN_FINDINGS, size=3, seed=7)
    second = rp.sample_rows(MAIN_FINDINGS, size=3, seed=7)
    assert first == second and len(first) == 3


def test_asking_for_more_than_there_is_takes_every_row():
    assert rp.sample_rows(MAIN_FINDINGS, size=99, seed=1) == MAIN_FINDINGS


def test_every_row_starts_with_no_mark_never_a_default():
    pack = rp.build_pack(MAIN_FINDINGS, SECONDARY_FINDINGS, main_sample_size=3, seed=0)
    assert pack["counts"] == {"main_available": 5, "main_sampled": 3, "secondary_sampled": 3}
    assert all(r["mark"] is None and r["note"] == "" for r in pack["rows"])
    assert {r["section"] for r in pack["rows"]} == {rp.MAIN, rp.SECONDARY}


def test_the_secondary_set_is_not_sampled_by_default_its_a_fixed_set():
    pack = rp.build_pack(MAIN_FINDINGS, SECONDARY_FINDINGS, main_sample_size=2, seed=0)
    assert pack["counts"]["secondary_sampled"] == len(SECONDARY_FINDINGS)


# ------------------------------------------------------------- the workbook

def test_rendering_and_reading_back_an_unfilled_pack_shows_no_marks(tmp_path):
    pack = rp.build_pack(MAIN_FINDINGS, SECONDARY_FINDINGS, main_sample_size=2, seed=0)
    path = tmp_path / "pack.xlsx"
    path.write_bytes(rp.render_pack(pack))

    filled = rp.read_pack(path)
    assert len(filled["rows"]) == len(pack["rows"])
    assert all(r["mark"] is None and not r.get("mark_error") for r in filled["rows"])
    assert filled["missed_defects"] == 0


def test_a_pack_with_no_rows_still_renders_and_reads_back(tmp_path):
    pack = rp.build_pack([], [], main_sample_size=5, seed=0)
    path = tmp_path / "empty.xlsx"
    path.write_bytes(rp.render_pack(pack))
    assert rp.read_pack(path)["rows"] == []


def test_a_sheet_with_the_wrong_columns_is_refused(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Review pack"
    wb.active.append(["Not", "The", "Right", "Columns"])
    path = tmp_path / "wrong.xlsx"
    wb.save(path)
    try:
        rp.read_pack(path)
        assert False, "should have refused"
    except rp.PackError as exc:
        assert "columns do not match" in str(exc)


def test_a_file_with_no_review_pack_sheet_is_refused(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Not", "A", "Pack"])
    path = tmp_path / "not-a-pack.xlsx"
    wb.save(path)
    try:
        rp.read_pack(path)
        assert False, "should have refused"
    except rp.PackError as exc:
        assert "no 'Review pack' sheet" in str(exc)


# ------------------------------------------------------------------ scoring

def _fill(path, marks: dict[int, str], missed: int = 0):
    """Load a rendered pack and write marks onto it by item number, plus
    `missed` invented missed-defect rows, exactly as an engineer would."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["Review pack"]
    col = rp.COLUMNS.index("Mark") + 1
    for row in ws.iter_rows(min_row=2):
        item_no = row[0].value
        if item_no in marks:
            row[col - 1].value = marks[item_no]
    missed_ws = wb[rp.MISSED_DEFECTS_SHEET]
    for n in range(missed):
        missed_ws.append([f"STD-{n}", "9.9", "an invented missed defect"])
    wb.save(path)


def test_the_score_counts_each_mark_apart_per_section_never_folded_together(tmp_path):
    pack = rp.build_pack(MAIN_FINDINGS, SECONDARY_FINDINGS, main_sample_size=3, seed=0)
    path = tmp_path / "pack.xlsx"
    path.write_bytes(rp.render_pack(pack))
    main_items = [r["item_no"] for r in pack["rows"] if r["section"] == rp.MAIN]
    secondary_items = [r["item_no"] for r in pack["rows"] if r["section"] == rp.SECONDARY]
    _fill(path, {main_items[0]: "right", main_items[1]: "wrong",
                secondary_items[0]: "unsure"}, missed=2)

    result = rp.score_pack(rp.read_pack(path))

    assert (result[rp.MAIN]["right"], result[rp.MAIN]["wrong"]) == (1, 1)
    assert result[rp.MAIN]["unmarked_item_nos"] == [main_items[2]]
    assert result[rp.SECONDARY]["unsure"] == 1
    assert result["missed_defects"] == 2
    assert cli.main(["score", "--pack", str(path)]) == 0


def test_an_unreadable_mark_is_reported_apart_never_counted_as_right_or_dropped(tmp_path):
    pack = rp.build_pack(MAIN_FINDINGS, [], main_sample_size=2, seed=0)
    path = tmp_path / "pack.xlsx"
    path.write_bytes(rp.render_pack(pack))
    item_no = pack["rows"][0]["item_no"]
    _fill(path, {item_no: "maybe-ish"})  # not one of right/wrong/unsure

    filled = rp.read_pack(path)
    result = rp.score_pack(filled)
    assert result[rp.MAIN]["unreadable_mark_item_nos"] == [item_no]
    assert sum(result[rp.MAIN][m] for m in rp.MARKS) == 0  # not silently counted


def test_an_unmarked_row_is_named_by_item_number_never_hidden():
    pack = rp.build_pack(MAIN_FINDINGS, [], main_sample_size=3, seed=0)
    filled = {"rows": [{"section": rp.MAIN, "item_no": r["item_no"], "mark": None,
                        "mark_error": None, "note": ""} for r in pack["rows"]],
             "missed_defects": 0}
    result = rp.score_pack(filled)
    assert sorted(result[rp.MAIN]["unmarked_item_nos"]) == sorted(r["item_no"] for r in pack["rows"])


def test_score_never_prints_a_findings_text_only_ids_and_counts(tmp_path, capsys):
    secret = "THE-CLIENTS-OWN-WORDS-MUST-NEVER-APPEAR-HERE"
    main = [_finding("m1", text=secret)]
    pack = rp.build_pack(main, [], main_sample_size=1, seed=0)
    path = tmp_path / "pack.xlsx"
    path.write_bytes(rp.render_pack(pack))
    _fill(path, {pack["rows"][0]["item_no"]: "right"})

    cli.main(["score", "--pack", str(path)])
    assert secret not in capsys.readouterr().out


# --------------------------------------------------------------------- CLI

def test_the_cli_refuses_to_write_the_pack_inside_the_repo(tmp_path):
    main_file = tmp_path / "main.findings.json"
    main_file.write_text(json.dumps({"format": "review-findings/1", "findings": MAIN_FINDINGS}))
    sec_file = tmp_path / "secondary.findings.json"
    sec_file.write_text(json.dumps({"format": "review-findings/1", "findings": SECONDARY_FINDINGS}))
    inside = REPO / "eval" / "review" / "_test_886_refused.xlsx"

    code = cli.main(["build", "--main", str(main_file), "--secondary", str(sec_file),
                     "--main-sample", "2", "--out", str(inside)])

    assert code == 2 and not inside.exists()


def test_the_cli_builds_a_pack_from_findings_files_on_disk(tmp_path, capsys):
    main_file = tmp_path / "main.findings.json"
    main_file.write_text(json.dumps({"format": "review-findings/1", "findings": MAIN_FINDINGS}))
    sec_file = tmp_path / "secondary.findings.json"
    sec_file.write_text(json.dumps({"format": "review-findings/1", "findings": SECONDARY_FINDINGS}))
    out = tmp_path / "pack.xlsx"

    code = cli.main(["build", "--main", str(main_file), "--secondary", str(sec_file),
                     "--main-sample", "2", "--out", str(out)])

    assert code == 0 and out.exists()
    printed = capsys.readouterr().out
    assert "2 MAIN row(s)" in printed and f"{len(SECONDARY_FINDINGS)} SECONDARY row(s)" in printed
